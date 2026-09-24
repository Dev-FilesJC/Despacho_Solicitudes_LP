import streamlit as st
import pandas as pd
from io import BytesIO
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo
import time
import re
import uuid
import hashlib

from streamlit_autorefresh import st_autorefresh

from supabase import create_client, Client

# =========================================================
# CONFIGURACION
# =========================================================
st.set_page_config(
    page_title="JC Control de Solicitudes — Despacho",
    page_icon="📝",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# =========================================================
# ESTILOS
# =========================================================
st.markdown(
    """
<style>
* { box-sizing:border-box; }
.stApp { background:#020914; color:#f3f4f6; font-family:Arial,Helvetica,sans-serif; }
.block-container { width:94%; max-width:1450px; padding-top:22px!important; padding-bottom:24px!important; margin:auto; }
header[data-testid="stHeader"] { background:transparent; }
div[data-testid="stToolbar"] { display:none; }
footer { display:none; }
h1,h2,h3 { color:#f3f4f6!important; font-family:Arial,Helvetica,sans-serif!important; }
.titulo { text-align:center; margin:4px 0 14px; font-size:30px; font-weight:900; line-height:1.2; width:100%; white-space:nowrap; }
.stButton>button,.stDownloadButton>button { border:1px solid #273246!important; border-radius:5px!important; min-height:30px!important; height:30px!important; padding:4px 8px!important; font-weight:800!important; color:white!important; background:#17263d!important; font-size:10px!important; }
.stButton>button:hover,.stDownloadButton>button:hover { filter:brightness(1.18); border-color:#315b86!important; }

/* BOTONES CANCELAR — rojo y respuesta visual al presionarlos */
.stFormSubmitButton > button[kind="primary"] {
    background:#7f1d1d!important;
    border-color:#ef4444!important;
    color:#ffffff!important;
}
.stFormSubmitButton > button[kind="primary"]:hover {
    background:#b91c1c!important;
    border-color:#f87171!important;
    filter:none!important;
}
.stFormSubmitButton > button[kind="primary"]:active {
    background:#ef4444!important;
    border-color:#fecaca!important;
    transform:scale(.98)!important;
}
div[data-testid="stForm"] { background:rgba(7,15,29,.55); border:1px solid #273246; border-radius:6px; padding:10px!important; }
div[data-testid="stForm"] label { color:#cbd5e1!important; font-size:10px!important; font-weight:700!important; }
div[data-baseweb="input"],div[data-baseweb="select"]>div,div[data-testid="stDateInput"]>div { background:#242630!important; color:#f3f4f6!important; border-radius:5px!important; border-color:transparent!important; min-height:33px!important; }
div[data-baseweb="input"] input,div[data-testid="stDateInput"] input { color:#f3f4f6!important; font-size:11px!important; }
div[data-baseweb="select"] span { color:#f3f4f6!important; font-size:11px!important; }
div[data-baseweb="select"] svg { fill:#f3f4f6!important; }
div[data-testid="stTextInput"] input { width:100%; height:33px; border:1px solid transparent; border-radius:5px; background:#242630!important; color:#f3f4f6!important; padding:0 9px; font-size:11px; }
/* =========================================================
   TABLA JC — ESTILO PROFESIONAL
   ========================================================= */
div[data-testid="stDataFrame"],
div[data-testid="stDataEditor"] {
    width:100%!important;
    border:1px solid #30445f!important;
    border-radius:12px!important;
    overflow:hidden!important;
    background:linear-gradient(180deg,#0b1526 0%,#07101d 100%)!important;
    box-shadow:0 8px 28px rgba(0,0,0,.28), inset 0 1px 0 rgba(255,255,255,.025)!important;
}

/* Barra superior del editor */
div[data-testid="stDataEditor"] > div {
    border-radius:12px!important;
}

/* Encabezados */
div[data-testid="stDataEditor"] [role="columnheader"] {
    background:#172a43!important;
    color:#f8fafc!important;
    font-weight:800!important;
    border-bottom:1px solid #3a5270!important;
}

/* Indicadores de color dentro de ESTADO y PRIORIDAD */
div[data-testid="stDataEditor"] [role="gridcell"] {
    font-size: 12px !important;
}

/* Encabezados de acciones */
div[data-testid="stDataEditor"] [role="columnheader"] {
    letter-spacing: .2px !important;
}

/* Botones de edición de la tabla */
div[data-testid="stDataEditor"] button {
    border-radius:6px!important;
}

/* Checkboxes EDITAR / ELIMINAR */
div[data-testid="stDataEditor"] input[type="checkbox"] {
    accent-color:#38bdf8!important;
}

/* Contenedor de la tabla */
div[data-testid="stDataEditor"] canvas {
    border-radius:10px!important;
}

/* Texto que acompaña el contador */
div[data-testid="stCaptionContainer"] {
    color:#94a3b8!important;
    font-size:11px!important;
    font-weight:700!important;
    padding:5px 2px 8px!important;
}

/* Separador */
hr { border-color:#273246!important; }
.footer { margin-top:30px; padding:15px 10px; border-top:1px solid #273246; text-align:center; color:#8f9bad; font-size:11px; }
.footer strong { color:#dbe2ea; }
.login-wrapper { max-width:380px; margin:8px auto 14px; text-align:center; }
.login-icon { font-size:30px; line-height:1; margin-bottom:5px; }
.login-title { font-size:27px; font-weight:800; line-height:1.15; }
.login-subtitle { font-size:12px; opacity:.62; margin-top:5px; }
</style>
""",
    unsafe_allow_html=True,
)

# =========================================================
# CONSTANTES
# =========================================================
COLS = [
    "CLIENTE", "NRO SOLICITUD - WO", "TIPO DE SOLICITUD", "CENTRO DE COSTO",
    "PRIORIDAD", "CANT - ITEMS", "ESTADO DE SOLICITUD", "DIRECCIÓN", "FECHA DE INGRESO"
]
RUTA_COLS = [
    "REGIONAL", "CENTRO DE ACOPIO", "AGENCIAS",
    "FECHA LIMITE DE INGRESO (SE1)",
    "FECHA LIMITE DE INGRESO (SR1)",
    "FECHA DE RECOJO"
]
CAJAS_COLS = [
    "SOLICITANTE", "CLIENTE", "AGENCIA", "FECHA DE SALIDA",
    "CANTIDAD DE CAJAS SOLICITADAS", "NRO. WORKORDER",
    "CAJAS REUTILIZADAS", "CANTIDAD DE CINTILLOS", "OBSERVACIONES"
]
COMISION_COLS = [
    "OPERADOR", "FECHA", "TIPO DE SALIDA", "HORA DE INGRESO",
    "HORA DE SALIDA", "TIEMPO ADICIONAL", "TOTAL HORAS", "COMPENSADO",
    "FECHA COMPENSACIÓN", "DESCRIPCIÓN"
]
RUTA_SHEET = "PROGRAMACION_RUTAS"
LOCK_TIMEOUT_SECONDS = 15 * 60
INACTIVITY_TIMEOUT_SECONDS = 8 * 60
AUTO_REFRESH_INTERVAL_MS = 60 * 1000
ZONA_HORARIA_APP = ZoneInfo("America/La_Paz")

PRIORIDADES = ["Seleccione una opcion", "RUSH", "TURNO SIGUIENTE", "NORMAL"]
REGIONALES = ["Seleccione una opcion", "LA PAZ", "COCHABAMBA", "SANTA CRUZ", "SUCRE", "ORURO", "POTOSI", "OTRA"]

# Los catálogos (clientes, tipos, centros, estados, direcciones, agencias,
# operadores y tipos de salida) se cargan directamente desde Supabase.

# =========================================================
# SESSION STATE
# =========================================================
def ss(name, default):
    if name not in st.session_state:
        st.session_state[name] = default

ss("rows", None)
ss("rutas", None)
ss("cajas", None)
ss("movimientos_stock", None)
ss("comisiones", None)
ss("comision_form_version", 0)
ss("mostrar_form_comision", False)
ss("comision_id_editar_inline", None)
ss("comision_tabla_version", 0)
ss("page", "solicitudes")
ss("caja_form_version", 0)
ss("stock_form_version", 0)
ss("mostrar_editor_historial_salidas", False)
ss("form_version", 0)
ss("tabla_version", 0)
ss("ruta_form_version", 0)
ss("editing", None)
ss("editing_key", "")
ss("session_id", uuid.uuid4().hex)
ss("registro_activo", False)
ss("ruta_activo", False)
ss("usuario_nombre", "")
ss("autenticado", False)
ss("busqueda_masiva_ids", [])
ss("busqueda_masiva_nombre", "")
ss("busqueda_masiva_version", 0)
ss("kardex_datos_editados", None)
ss("kardex_clientes_editados", None)
ss("kardex_editor_firma", "")

# =========================================================
# SUPABASE
# =========================================================
@st.cache_resource
def obtener_supabase() -> Client:
    try:
        url = st.secrets["SUPABASE_URL"]
        key = (
            st.secrets.get("SUPABASE_SECRET_KEY")
            or st.secrets.get("SUPABASE_SERVICE_ROLE_KEY")
            or st.secrets.get("SUPABASE_KEY")
        )
    except Exception as e:
        raise RuntimeError("Configura SUPABASE_URL y SUPABASE_KEY en Streamlit Secrets.") from e
    if not url or not key:
        raise RuntimeError("Faltan SUPABASE_URL y/o SUPABASE_KEY en Streamlit Secrets.")
    return create_client(url, key)


def db():
    return obtener_supabase()


@st.cache_data(ttl=120, show_spinner=False)
def cargar_catalogo(categoria, incluir_placeholder=True):
    """Carga un catálogo activo desde Supabase y lo mantiene en caché por 2 minutos."""
    response = (
        db().table("catalogos_app")
        .select("valor")
        .eq("categoria", categoria)
        .eq("activo", True)
        .order("valor", desc=False)
        .execute()
    )
    valores = [str(r.get("valor", "")).strip() for r in (response.data or [])]
    valores = [v for v in valores if v]
    return (["Seleccione una opcion"] + valores) if incluir_placeholder else valores


def agregar_catalogo(categoria, valor):
    """Agrega o reactiva un valor de catálogo."""
    valor = str(valor or "").strip()
    if not valor:
        raise ValueError("El valor no puede estar vacío.")
    existente = (
        db().table("catalogos_app")
        .select("id, activo")
        .eq("categoria", categoria)
        .eq("valor", valor)
        .limit(1)
        .execute()
    )
    if existente.data:
        respuesta = db().table("catalogos_app").update({"activo": True}).eq("id", existente.data[0]["id"]).execute()
    else:
        respuesta = db().table("catalogos_app").insert({"categoria": categoria, "valor": valor, "activo": True}).execute()
    cargar_catalogo.clear()
    return respuesta


def desactivar_catalogo(categoria, valor):
    """Desactiva un valor sin borrar las solicitudes/registros históricos."""
    respuesta = (
        db().table("catalogos_app")
        .update({"activo": False})
        .eq("categoria", categoria)
        .eq("valor", valor)
        .execute()
    )
    cargar_catalogo.clear()
    return respuesta


def fecha_local_hoy():
    return datetime.now(ZONA_HORARIA_APP).date()


def convertir_fecha(valor):
    """Convierte ISO o DD/MM/YYYY sin invertir mes y día."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    if not texto:
        return None
    for formato in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:26], formato).date()
        except ValueError:
            pass
    convertido = pd.to_datetime(texto, errors="coerce", dayfirst=False)
    return None if pd.isna(convertido) else convertido.date()


def normalizar(v):
    return str(v if pd.notna(v) else "").strip().upper()


def extraer_solicitudes_de_texto(texto):
    """Extrae números de solicitud/WO separados por líneas, comas o punto y coma."""
    if not texto:
        return []
    partes = re.split(r"[\n,;]+", str(texto))
    valores = []
    for parte in partes:
        valor = normalizar(parte)
        if valor and valor not in valores:
            valores.append(valor)
    return valores


def extraer_solicitudes_de_archivo(uploaded_file):
    """Lee TXT/CSV/XLSX y devuelve todos los valores de sus celdas.
    La coincidencia final se hace contra los NRO SOLICITUD - WO existentes,
    por lo que encabezados u otros datos del archivo no generan falsos resultados.
    """
    nombre = str(getattr(uploaded_file, "name", "")).lower()
    if nombre.endswith(".txt"):
        contenido = uploaded_file.getvalue().decode("utf-8-sig", errors="ignore")
        return extraer_solicitudes_de_texto(contenido)

    if nombre.endswith(".csv"):
        try:
            datos = pd.read_csv(uploaded_file, header=None, dtype=str, keep_default_na=False)
        except Exception:
            uploaded_file.seek(0)
            datos = pd.read_csv(uploaded_file, header=None, dtype=str, keep_default_na=False, sep=";")
        return [normalizar(v) for v in datos.to_numpy().ravel().tolist() if normalizar(v)]

    if nombre.endswith(".xlsx"):
        libro = pd.ExcelFile(uploaded_file, engine="openpyxl")
        candidatos = []
        for hoja in libro.sheet_names:
            datos = pd.read_excel(libro, sheet_name=hoja, header=None, dtype=str)
            candidatos.extend(datos.to_numpy().ravel().tolist())
        resultado = []
        for valor in candidatos:
            valor_norm = normalizar(valor)
            if valor_norm and valor_norm not in resultado:
                resultado.append(valor_norm)
        return resultado

    raise ValueError("Formato no compatible. Usa .xlsx, .csv o .txt.")

# =========================================================
# SOLICITUDES
# =========================================================
def supabase_select_all(table_name, columns="*", orderings=None, page_size=1000, filters=None):
    """Lee todos los registros de una tabla de Supabase por bloques.

    Supabase/PostgREST puede limitar una consulta a 1.000 filas.
    Este helper pagina automáticamente para que la aplicación no deje
    de mostrar/buscar registros al superar ese límite.
    """
    orderings = orderings or []
    filters = filters or []
    all_data = []
    offset = 0

    while True:
        query = db().table(table_name).select(columns)
        for method, column, value in filters:
            query = getattr(query, method)(column, value)
        for column, descending in orderings:
            query = query.order(column, desc=descending)

        response = query.range(offset, offset + page_size - 1).execute()
        chunk = response.data or []
        all_data.extend(chunk)

        if len(chunk) < page_size:
            break
        offset += page_size

    return all_data


def cargar_solicitudes_supabase():
    data = supabase_select_all(
        "solicitudes",
        "id, cliente, nro_solicitud_wo, tipo_solicitud, centro_costo, prioridad, cantidad_items, estado_solicitud, direccion, fecha_ingreso",
        orderings=[("id", True)],
    )
    rows = []
    for r in data:
        rows.append({
            "_ID_": r.get("id"),
            "CLIENTE": r.get("cliente", ""),
            "NRO SOLICITUD - WO": r.get("nro_solicitud_wo", ""),
            "TIPO DE SOLICITUD": r.get("tipo_solicitud", ""),
            "CENTRO DE COSTO": r.get("centro_costo", ""),
            "PRIORIDAD": r.get("prioridad", ""),
            "CANT - ITEMS": r.get("cantidad_items", 0),
            "ESTADO DE SOLICITUD": r.get("estado_solicitud", ""),
            "DIRECCIÓN": r.get("direccion", ""),
            "FECHA DE INGRESO": r.get("fecha_ingreso", ""),
        })
    return pd.DataFrame(rows, columns=["_ID_"] + COLS).fillna("")


def guardar_solicitud_supabase(nuevo, registro_id=None):
    payload = {
        "cliente": nuevo["CLIENTE"],
        "nro_solicitud_wo": nuevo["NRO SOLICITUD - WO"],
        "tipo_solicitud": nuevo["TIPO DE SOLICITUD"],
        "centro_costo": nuevo.get("CENTRO DE COSTO", ""),
        "prioridad": nuevo["PRIORIDAD"],
        "cantidad_items": int(nuevo["CANT - ITEMS"]),
        "estado_solicitud": nuevo["ESTADO DE SOLICITUD"],
        "direccion": nuevo["DIRECCIÓN"],
        "fecha_ingreso": nuevo["FECHA DE INGRESO"],
    }
    q = db().table("solicitudes")
    return q.update(payload).eq("id", int(registro_id)).execute() if registro_id else q.insert(payload).execute()


def eliminar_solicitudes_supabase(ids):
    for registro_id in ids:
        db().table("solicitudes").delete().eq("id", int(registro_id)).execute()

def cargar_solicitudes_mes_supabase(fecha_inicio, fecha_fin_exclusiva):
    data = supabase_select_all(
        "solicitudes",
        "*",
        orderings=[("fecha_ingreso", True), ("id", True)],
        filters=[
            ("gte", "fecha_ingreso", fecha_inicio.isoformat()),
            ("lt", "fecha_ingreso", fecha_fin_exclusiva.isoformat()),
        ],
    )
    rows = []
    for r in data:
        rows.append({
            "_ID_": r.get("id"),
            "CLIENTE": r.get("cliente", ""),
            "NRO SOLICITUD - WO": r.get("nro_solicitud_wo", ""),
            "TIPO DE SOLICITUD": r.get("tipo_solicitud", ""),
            "CENTRO DE COSTO": r.get("centro_costo", ""),
            "PRIORIDAD": r.get("prioridad", ""),
            "CANT - ITEMS": r.get("cantidad_items", 0),
            "ESTADO DE SOLICITUD": r.get("estado_solicitud", ""),
            "DIRECCIÓN": r.get("direccion", ""),
            "FECHA DE INGRESO": r.get("fecha_ingreso", ""),
        })
    return pd.DataFrame(rows, columns=["_ID_"] + COLS).fillna("")

# =========================================================
# PROGRAMACION DE RUTAS
# La tabla programacion_rutas usa columnas directas:
# id, regional, centro_acopio, agencias,
# fecha_limite_ingreso_se1, fecha_limite_ingreso_sr1, fecha_recojo.
# =========================================================
def cargar_rutas_supabase():
    data = supabase_select_all(
        "programacion_rutas",
        "id, regional, centro_acopio, agencias, fecha_limite_ingreso_se1, fecha_limite_ingreso_sr1, fecha_recojo",
        orderings=[("id", False)],
    )
    rows = []
    for r in data:
        rows.append({
            "_ID_RUTA_": r.get("id"),
            "REGIONAL": r.get("regional", ""),
            "CENTRO DE ACOPIO": r.get("centro_acopio", ""),
            "AGENCIAS": r.get("agencias", ""),
            "FECHA LIMITE DE INGRESO (SE1)": r.get("fecha_limite_ingreso_se1", ""),
            "FECHA LIMITE DE INGRESO (SR1)": r.get("fecha_limite_ingreso_sr1", ""),
            "FECHA DE RECOJO": r.get("fecha_recojo", ""),
        })
    return pd.DataFrame(rows, columns=["_ID_RUTA_"] + RUTA_COLS).fillna("")


def guardar_ruta_supabase(datos, ruta_id=None):
    payload = {
        "regional": datos["REGIONAL"],
        "centro_acopio": datos["CENTRO DE ACOPIO"],
        "agencias": datos["AGENCIAS"],
        "fecha_limite_ingreso_se1": datos["FECHA LIMITE DE INGRESO (SE1)"],
        "fecha_limite_ingreso_sr1": datos["FECHA LIMITE DE INGRESO (SR1)"],
        "fecha_recojo": datos["FECHA DE RECOJO"],
    }
    q = db().table("programacion_rutas")
    if ruta_id:
        return q.update(payload).eq("id", int(ruta_id)).execute()
    return q.insert(payload).execute()


def eliminar_ruta_supabase(ruta_id):
    return db().table("programacion_rutas").delete().eq("id", int(ruta_id)).execute()


# =========================================================
# CONTROL DE CAJAS Y CINTILLOS
# =========================================================
def cargar_cajas_supabase():
    # Igual que Control de Salidas a Comisión: primero las salidas más recientes.
    data = supabase_select_all(
        "control_cajas",
        "*",
        orderings=[("fecha_salida", True), ("id", True)],
    )
    rows = []
    for r in data:
        rows.append({
            "_ID_CAJA_": r.get("id"),
            "SOLICITANTE": r.get("solicitante", ""),
            "CLIENTE": r.get("cliente", ""),
            "AGENCIA": r.get("agencia", ""),
            "FECHA DE SALIDA": r.get("fecha_salida", ""),
            "CANTIDAD DE CAJAS SOLICITADAS": r.get("cantidad_cajas_solicitadas", 0),
            "NRO. WORKORDER": r.get("nro_workorder", ""),
            "CAJAS REUTILIZADAS": r.get("cajas_reutilizadas", 0),
            "CANTIDAD DE CINTILLOS": r.get("cantidad_cintillos", 0),
            "OBSERVACIONES": r.get("observaciones", ""),
        })
    return pd.DataFrame(rows, columns=["_ID_CAJA_"] + CAJAS_COLS).fillna("")


def cargar_movimientos_stock_supabase():
    data = supabase_select_all(
        "movimientos_stock",
        "*",
        orderings=[("fecha", True), ("id", True)],
    )
    return pd.DataFrame(data)


def guardar_movimiento_stock(fecha, tipo, cajas_nuevas=0, cajas_reutilizadas=0, cintillos=0, proveedor="", observaciones=""):
    payload = {
        "fecha": fecha.isoformat() if isinstance(fecha, date) else str(fecha),
        "tipo_movimiento": tipo,
        "cajas_nuevas": int(cajas_nuevas or 0),
        "cajas_reutilizadas": int(cajas_reutilizadas or 0),
        "cintillos": int(cintillos or 0),
        "proveedor": (proveedor or "").strip(),
        "observaciones": (observaciones or "").strip(),
        "usuario": st.session_state.get("usuario_nombre", ""),
    }
    return db().table("movimientos_stock").insert(payload).execute()


def obtener_stock_actual():
    mov = cargar_movimientos_stock_supabase()
    if mov.empty:
        return {"CAJAS NUEVAS": 0, "CAJAS REUTILIZADAS": 0, "CINTILLOS": 0}
    return {
        "CAJAS NUEVAS": int(pd.to_numeric(mov.get("cajas_nuevas", 0), errors="coerce").fillna(0).sum()),
        "CAJAS REUTILIZADAS": int(pd.to_numeric(mov.get("cajas_reutilizadas", 0), errors="coerce").fillna(0).sum()),
        "CINTILLOS": int(pd.to_numeric(mov.get("cintillos", 0), errors="coerce").fillna(0).sum()),
    }


def guardar_salida_cajas_supabase(datos):
    payload = {
        "solicitante": datos["SOLICITANTE"],
        "cliente": datos["CLIENTE"],
        "agencia": datos["AGENCIA"],
        "fecha_salida": datos["FECHA DE SALIDA"],
        "cantidad_cajas_solicitadas": int(datos["CANTIDAD DE CAJAS SOLICITADAS"]),
        "nro_workorder": datos["NRO. WORKORDER"],
        "cajas_reutilizadas": int(datos["CAJAS REUTILIZADAS"]),
        "cantidad_cintillos": int(datos["CANTIDAD DE CINTILLOS"]),
        "observaciones": datos["OBSERVACIONES"],
        "usuario": st.session_state.get("usuario_nombre", ""),
    }
    return db().table("control_cajas").insert(payload).execute()


def entero_seguro(valor, default=0):
    """Convierte valores de tabla a entero sin fallar con vacíos o NaN."""
    try:
        numero = pd.to_numeric(valor, errors="coerce")
        if pd.isna(numero):
            return int(default)
        return int(numero)
    except Exception:
        return int(default)


def eliminar_salida_cajas_supabase(salida_id):
    """Elimina un registro del historial de salidas."""
    return db().table("control_cajas").delete().eq("id", int(salida_id)).execute()


def actualizar_salida_cajas_supabase(datos, salida_id):
    """Actualiza un registro existente del historial de salidas."""
    payload = {
        "solicitante": datos["SOLICITANTE"],
        "cliente": datos["CLIENTE"],
        "agencia": datos["AGENCIA"],
        "fecha_salida": datos["FECHA DE SALIDA"],
        "cantidad_cajas_solicitadas": int(datos["CANTIDAD DE CAJAS SOLICITADAS"]),
        "nro_workorder": datos["NRO. WORKORDER"],
        "cajas_reutilizadas": int(datos["CAJAS REUTILIZADAS"]),
        "cantidad_cintillos": int(datos["CANTIDAD DE CINTILLOS"]),
        "observaciones": datos["OBSERVACIONES"],
        "usuario": st.session_state.get("usuario_nombre", ""),
    }
    return db().table("control_cajas").update(payload).eq("id", int(salida_id)).execute()

# =========================================================
# CONTROL DE SALIDAS A COMISIÓN
# =========================================================
def hora_actual_local():
    """Hora actual de La Paz, redondeada al minuto para registrar en el instante."""
    return datetime.now(ZONA_HORARIA_APP).time().replace(second=0, microsecond=0)


def calcular_total_horas(fecha_comision, hora_ingreso, hora_salida):
    """Calcula horas entre ingreso y salida, permitiendo pasar medianoche."""
    if hora_ingreso is None or hora_salida is None:
        return None
    total = (
        datetime.combine(fecha_comision, hora_salida)
        - datetime.combine(fecha_comision, hora_ingreso)
    ).total_seconds() / 3600
    if total < 0:
        total += 24
    return round(total, 4)


def formatear_horas_minutos(horas):
    """Muestra horas decimales de forma clara: 1.50 -> 1 h 30 min."""
    try:
        minutos_totales = int(round(float(horas) * 60))
    except (TypeError, ValueError):
        return "0 h 00 min"

    horas_enteras, minutos = divmod(minutos_totales, 60)
    return f"{horas_enteras} h {minutos:02d} min"


TIEMPOS_ADICIONALES = {
    "Sin aumento": 0.0,
    "30 minutos": 0.5,
    "1 hora": 1.0,
    "1 hora 30 minutos": 1.5,
    "2 horas": 2.0,
    "2 horas 30 minutos": 2.5,
    "3 horas": 3.0,
    "4 horas": 4.0,
    "5 horas": 5.0,
}


def etiqueta_tiempo_adicional(horas):
    """Convierte horas adicionales almacenadas en una etiqueta amigable."""
    try:
        valor = float(horas or 0)
    except (TypeError, ValueError):
        valor = 0.0
    for etiqueta, cantidad in TIEMPOS_ADICIONALES.items():
        if abs(valor - cantidad) < 0.001:
            return etiqueta
    return formatear_horas_minutos(valor)


def es_valor_compensado(valor):
    """Interpreta correctamente booleanos de Supabase y textos antiguos."""
    if isinstance(valor, str):
        return valor.strip().lower() in ("true", "1", "si", "sí", "yes")
    return bool(valor)


def cargar_comisiones_supabase():
    data = supabase_select_all(
        "control_comisiones",
        "*",
        orderings=[("fecha", True), ("id", True)],
    )
    rows = []
    for r in data:
        rows.append({
            "_ID_COMISION_": r.get("id"),
            "OPERADOR": r.get("operador", ""),
            "FECHA": r.get("fecha", ""),
            "TIPO DE SALIDA": r.get("tipo_salida", ""),
            "HORA DE INGRESO": r.get("hora_ingreso", ""),
            "HORA DE SALIDA": r.get("hora_salida", ""),
            "TIEMPO ADICIONAL": r.get("horas_adicionales", 0) or 0,
            "TOTAL HORAS": r.get("total_horas", ""),
            "COMPENSADO": es_valor_compensado(r.get("compensado", False)),
            "FECHA COMPENSACIÓN": r.get("fecha_compensacion", ""),
            "DESCRIPCIÓN": r.get("descripcion", ""),
        })
    return pd.DataFrame(rows, columns=["_ID_COMISION_"] + COMISION_COLS).fillna("")


def guardar_comision_supabase(datos):
    payload = {
        "operador": datos["OPERADOR"],
        "fecha": datos["FECHA"],
        "tipo_salida": datos["TIPO DE SALIDA"],
        "hora_ingreso": datos.get("HORA DE INGRESO") or None,
        "hora_salida": datos.get("HORA DE SALIDA") or None,
        "horas_adicionales": float(datos.get("TIEMPO ADICIONAL", 0) or 0),
        "total_horas": float(datos["TOTAL HORAS"]) if datos.get("TOTAL HORAS") not in (None, "") else None,
        "compensado": bool(datos.get("COMPENSADO", False)),
        "fecha_compensacion": datos.get("FECHA COMPENSACIÓN") or None,
        "descripcion": datos["DESCRIPCIÓN"],
        "usuario": st.session_state.get("usuario_nombre", ""),
    }
    return db().table("control_comisiones").insert(payload).execute()


def actualizar_comision_supabase(datos, comision_id):
    payload = {
        "operador": datos["OPERADOR"],
        "fecha": datos["FECHA"],
        "tipo_salida": datos["TIPO DE SALIDA"],
        "hora_ingreso": datos.get("HORA DE INGRESO") or None,
        "hora_salida": datos.get("HORA DE SALIDA") or None,
        "horas_adicionales": float(datos.get("TIEMPO ADICIONAL", 0) or 0),
        "total_horas": float(datos["TOTAL HORAS"]) if datos.get("TOTAL HORAS") not in (None, "") else None,
        "compensado": bool(datos.get("COMPENSADO", False)),
        "fecha_compensacion": datos.get("FECHA COMPENSACIÓN") or None,
        "descripcion": datos["DESCRIPCIÓN"],
        "usuario": st.session_state.get("usuario_nombre", ""),
    }
    return db().table("control_comisiones").update(payload).eq("id", int(comision_id)).execute()


def eliminar_comision_supabase(comision_id):
    """Elimina una salida a comisión por su ID."""
    return (
        db().table("control_comisiones")
        .delete()
        .eq("id", int(comision_id))
        .execute()
    )


def excel_bytes_comisiones(df):
    """Descarga exactamente el dataframe filtrado que se está mostrando."""
    out = BytesIO()
    datos = df.drop(columns=["_ID_COMISION_"], errors="ignore").copy()

    # Presentación amigable en Excel.
    if "COMPENSADO" in datos.columns:
        total_excel = pd.to_numeric(datos.get("TOTAL HORAS"), errors="coerce")
        datos["COMPENSADO"] = [
            "" if pd.isna(total)
            else ("SÍ" if es_valor_compensado(comp) else "NO")
            for total, comp in zip(total_excel, datos["COMPENSADO"])
        ]
    if "TIEMPO ADICIONAL" in datos.columns:
        datos["TIEMPO ADICIONAL"] = datos["TIEMPO ADICIONAL"].apply(
            lambda x: etiqueta_tiempo_adicional(x) if pd.notna(pd.to_numeric(x, errors="coerce")) else ""
        )
    if "TOTAL HORAS" in datos.columns:
        datos["TOTAL HORAS"] = pd.to_numeric(
            datos["TOTAL HORAS"], errors="coerce"
        ).map(lambda x: formatear_horas_minutos(x) if pd.notna(x) else "")

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        datos.to_excel(writer, sheet_name="SALIDAS COMISION", index=False)

        # Resumen de horas del mismo filtro descargado.
        total = pd.to_numeric(datos.get("TOTAL HORAS", pd.Series(dtype=float)), errors="coerce").fillna(0)
        compensado_mask = datos.get("COMPENSADO", pd.Series(dtype=object)).astype(str).str.upper().eq("SÍ")
        horas_compensadas = float(total[compensado_mask].sum()) if len(total) else 0.0
        horas_pendientes = float(total.sum()) - horas_compensadas

        resumen = pd.DataFrame({
            "INDICADOR": [
                "REGISTROS FILTRADOS",
                "HORAS ACUMULADAS",
                "HORAS COMPENSADAS",
                "HORAS PENDIENTES DE COMPENSAR",
            ],
            "VALOR": [
                len(datos),
                formatear_horas_minutos(float(total.sum())),
                formatear_horas_minutos(horas_compensadas),
                formatear_horas_minutos(horas_pendientes),
            ],
        })
        resumen.to_excel(writer, sheet_name="RESUMEN HORAS", index=False)

        # Fechas reales de Excel con formato DD/MM/YYYY.
        formatear_fechas_excel(
            writer.sheets["SALIDAS COMISION"],
            ["FECHA", "FECHA COMPENSACIÓN"]
        )

        ws = writer.sheets["SALIDAS COMISION"]
        for col_cells in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)

        ws2 = writer.sheets["RESUMEN HORAS"]
        for col_cells in ws2.columns:
            max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
            ws2.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 18), 35)

    return out.getvalue()


# =========================================================
# BLOQUEO
# =========================================================
def _ahora_utc_iso():
    """Devuelve una fecha/hora UTC compatible con PostgreSQL timestamptz."""
    return datetime.now(timezone.utc).isoformat()


def _bloqueo_expirado(valor):
    """Comprueba si last_activity (timestamptz o valor antiguo Unix) expiró."""
    if valor in (None, "", 0, "0"):
        return True

    try:
        # Compatibilidad con registros antiguos que pudieron guardar time.time().
        if isinstance(valor, (int, float)):
            ultima = datetime.fromtimestamp(float(valor), tz=timezone.utc)
        else:
            texto = str(valor).strip()
            try:
                numero = float(texto)
                ultima = datetime.fromtimestamp(numero, tz=timezone.utc)
            except ValueError:
                ultima = pd.to_datetime(texto, utc=True, errors="coerce")
                if pd.isna(ultima):
                    return True
                ultima = ultima.to_pydatetime()

        if ultima.tzinfo is None:
            ultima = ultima.replace(tzinfo=timezone.utc)

        return (datetime.now(timezone.utc) - ultima).total_seconds() > LOCK_TIMEOUT_SECONDS
    except Exception:
        return True


def leer_bloqueo_edicion():
    try:
        response = db().table("app_locks").select("*").eq("id", 1).maybe_single().execute()
        row = response.data
        if not row or not row.get("owner_id"):
            return None

        if _bloqueo_expirado(row.get("last_activity")):
            liberar_bloqueo_edicion(force=True)
            return None

        return row
    except Exception:
        return None


def adquirir_bloqueo_edicion(motivo="registro"):
    usuario = (st.session_state.get("usuario_nombre") or "Usuario").strip()
    session_id = st.session_state.session_id
    actual = leer_bloqueo_edicion()

    if actual and actual.get("owner_id") != session_id:
        return False, actual

    # IMPORTANTE:
    # app_locks.last_activity es timestamp with time zone (timestamptz).
    # No debemos enviar time.time(), porque devuelve un Unix timestamp
    # como 1788320350.168386 y PostgreSQL lo rechaza como fecha.
    payload = {
        "id": 1,
        # app_locks.recurso es NOT NULL en Supabase.
        # Usamos el motivo de la operación como recurso para que
        # nunca se envíe NULL a esa columna.
        "recurso": str(motivo or "registro").strip() or "registro",
        "owner_id": session_id,
        "usuario": usuario,
        "last_activity": _ahora_utc_iso(),
        "motivo": str(motivo or "registro").strip() or "registro",
    }

    db().table("app_locks").upsert(payload, on_conflict="id").execute()
    st.session_state.registro_activo = True
    return True, payload


def renovar_bloqueo_edicion():
    if not st.session_state.get("registro_activo") and not st.session_state.get("ruta_activo"):
        return

    try:
        db().table("app_locks").update(
            {"last_activity": _ahora_utc_iso()}
        ).eq("id", 1).eq("owner_id", st.session_state.session_id).execute()
    except Exception:
        pass


def liberar_bloqueo_edicion(force=False):
    try:
        q = db().table("app_locks").delete().eq("id", 1)
        if not force:
            q = q.eq("owner_id", st.session_state.session_id)
        q.execute()
    except Exception:
        pass
    st.session_state.registro_activo = False
    st.session_state.ruta_activo = False



def formatear_fechas_excel(ws, nombres_columnas):
    """Convierte columnas de fecha a fechas reales de Excel y las muestra DD/MM/YYYY."""
    encabezados = {str(c.value).strip(): c.column for c in ws[1]}
    for nombre in nombres_columnas:
        col = encabezados.get(nombre)
        if not col:
            continue
        for fila in range(2, ws.max_row + 1):
            celda = ws.cell(fila, col)
            if celda.value in (None, ""):
                continue
            try:
                valor = celda.value
                if isinstance(valor, datetime):
                    fecha = valor.date()
                elif isinstance(valor, date):
                    fecha = valor
                else:
                    fecha = convertir_fecha(valor)
                if fecha is not None:
                    celda.value = fecha
                    celda.number_format = "DD/MM/YYYY"
            except Exception:
                pass

# =========================================================
# EXPORTAR EXCEL
# =========================================================

def preparar_editor_kardex_clientes(df):
    """Prepara una sola fila por CLIENTE para completar los datos del Kardex.

    El COD. CENTRO COSTO, los precios de BASE/TAPA y el TIPO DE CAMBIO
    se aplican automáticamente a todos los registros de ese cliente.
    """
    columnas = ["CLIENTE", "COD. CENTRO COSTO", "PRECIO BASE", "PRECIO TAPA", "TIPO DE CAMBIO"]
    if df is None or df.empty or "CLIENTE" not in df.columns:
        return pd.DataFrame(columns=columnas)

    clientes = (
        df["CLIENTE"]
        .fillna("")
        .astype(str)
        .str.strip()
    )
    clientes = sorted([x for x in clientes.unique().tolist() if x])
    base = pd.DataFrame({"CLIENTE": clientes})
    base["COD. CENTRO COSTO"] = ""
    base["PRECIO BASE"] = 0.0
    base["PRECIO TAPA"] = 0.0
    base["TIPO DE CAMBIO"] = 1.0
    return base[columnas]


def construir_kardex_por_registro(df_salidas, datos_clientes):
    """Expande la configuración por cliente a cada registro de salida."""
    columnas = [
        "_ID_CAJA_", "FECHA DE SALIDA", "NRO. WORKORDER", "CLIENTE", "AGENCIA",
        "CANTIDAD DE CAJAS SOLICITADAS", "COD. CENTRO COSTO", "PRECIO BASE",
        "PRECIO TAPA", "TIPO DE CAMBIO", "TOTAL BASE", "TOTAL TAPA"
    ]
    if df_salidas is None or df_salidas.empty:
        return pd.DataFrame(columns=columnas)

    config = {}
    if datos_clientes is not None and not datos_clientes.empty:
        for _, d in datos_clientes.iterrows():
            cliente = str(d.get("CLIENTE", "") or "").strip()
            if not cliente:
                continue
            config[cliente.upper()] = {
                "cod_centro": str(d.get("COD. CENTRO COSTO", "") or "").strip(),
                "precio_base": float(pd.to_numeric(d.get("PRECIO BASE", 0), errors="coerce") or 0),
                "precio_tapa": float(pd.to_numeric(d.get("PRECIO TAPA", 0), errors="coerce") or 0),
                "tipo_cambio": float(pd.to_numeric(d.get("TIPO DE CAMBIO", 1), errors="coerce") or 1),
            }

    filas = []
    for _, r in df_salidas.iterrows():
        cliente = str(r.get("CLIENTE", "") or "").strip()
        cfg = config.get(cliente.upper(), {
            "cod_centro": "", "precio_base": 0.0, "precio_tapa": 0.0, "tipo_cambio": 1.0
        })
        cajas = entero_seguro(r.get("CANTIDAD DE CAJAS SOLICITADAS"))
        filas.append({
            "_ID_CAJA_": r.get("_ID_CAJA_"),
            "FECHA DE SALIDA": r.get("FECHA DE SALIDA", ""),
            "NRO. WORKORDER": r.get("NRO. WORKORDER", ""),
            "CLIENTE": cliente,
            "AGENCIA": r.get("AGENCIA", ""),
            "CANTIDAD DE CAJAS SOLICITADAS": cajas,
            "COD. CENTRO COSTO": cfg["cod_centro"],
            "PRECIO BASE": cfg["precio_base"],
            "PRECIO TAPA": cfg["precio_tapa"],
            "TIPO DE CAMBIO": cfg["tipo_cambio"],
            "TOTAL BASE": round(cajas * cfg["precio_base"], 2),
            "TOTAL TAPA": round(cajas * cfg["precio_tapa"], 2),
        })
    return pd.DataFrame(filas, columns=columnas)


def resumen_tipo_solicitud_wo(df):
    """Cuenta SE1, SE2, SE3, SR1 y SR2 en la columna NRO SOLICITUD - WO."""
    tipos = ["SE1", "SE2", "SE3", "SR1", "SR2"]
    conteo = {tipo: 0 for tipo in tipos}

    if df is None or df.empty or "NRO SOLICITUD - WO" not in df.columns:
        return conteo

    serie = df["NRO SOLICITUD - WO"].fillna("").astype(str).str.upper()
    for tipo in tipos:
        conteo[tipo] = int(
            serie.apply(
                lambda valor: len(re.findall(rf"(?<![A-Z0-9]){tipo}(?![A-Z0-9])", valor))
            ).sum()
        )
    return conteo


def excel_bytes_mensual(df):
    """Genera el respaldo mensual con detalle completo y hojas de resumen."""
    out = BytesIO()
    datos = df.drop(columns=["_ID_"], errors="ignore").copy()

    # Asegurar que NRO SOLICITUD - WO esté incluido y visible.
    columnas_preferidas = [
        "NRO SOLICITUD - WO", "CLIENTE", "TIPO DE SOLICITUD",
        "CENTRO DE COSTO", "PRIORIDAD", "CANT - ITEMS",
        "ESTADO DE SOLICITUD", "DIRECCIÓN", "FECHA DE INGRESO"
    ]
    columnas = [c for c in columnas_preferidas if c in datos.columns]
    resto = [c for c in datos.columns if c not in columnas]
    datos = datos[columnas + resto]

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        datos.to_excel(writer, sheet_name="SOLICITUDES", index=False)

        # Resumen por estado
        if "ESTADO DE SOLICITUD" in datos.columns:
            resumen_estado = (
                datos["ESTADO DE SOLICITUD"]
                .fillna("SIN ESTADO")
                .astype(str)
                .value_counts()
                .rename_axis("ESTADO DE SOLICITUD")
                .reset_index(name="CANTIDAD")
            )
        else:
            resumen_estado = pd.DataFrame(columns=["ESTADO DE SOLICITUD", "CANTIDAD"])
        resumen_estado.to_excel(writer, sheet_name="RESUMEN_ESTADO", index=False)

        # Resumen por prioridad
        if "PRIORIDAD" in datos.columns:
            resumen_prioridad = (
                datos["PRIORIDAD"]
                .fillna("SIN PRIORIDAD")
                .astype(str)
                .value_counts()
                .rename_axis("PRIORIDAD")
                .reset_index(name="CANTIDAD")
            )
        else:
            resumen_prioridad = pd.DataFrame(columns=["PRIORIDAD", "CANTIDAD"])
        resumen_prioridad.to_excel(writer, sheet_name="RESUMEN_PRIORIDAD", index=False)

        # Resumen de tipos SE/SR desde NRO SOLICITUD - WO.
        # Se construye directamente para evitar referencias a variables locales
        # que puedan quedar sin inicializar al cambiar de mes.
        conteo_wo = resumen_tipo_solicitud_wo(datos)
        pd.DataFrame(
            {
                "TIPO WO": list(conteo_wo.keys()),
                "CANTIDAD": list(conteo_wo.values()),
            }
        ).to_excel(writer, sheet_name="RESUMEN_SE_SR", index=False)

        # Resumen general, incluyendo cantidad de solicitudes/WO.
        resumen_general = pd.DataFrame({
            "INDICADOR": [
                "TOTAL DE REGISTROS",
                "TOTAL DE SOLICITUDES / WO",
                "TOTAL DE ITEMS",
            ],
            "VALOR": [
                len(datos),
                datos["NRO SOLICITUD - WO"].nunique(dropna=True) if "NRO SOLICITUD - WO" in datos.columns else 0,
                pd.to_numeric(datos["CANT - ITEMS"], errors="coerce").fillna(0).sum()
                if "CANT - ITEMS" in datos.columns else 0,
            ]
        })
        resumen_general.to_excel(writer, sheet_name="RESUMEN_GENERAL", index=False)

        # Fechas reales de Excel, visibles como DD/MM/YYYY.
        formatear_fechas_excel(writer.sheets["SOLICITUDES"], ["FECHA DE INGRESO"])

        # Ajustar ancho de columnas para facilitar la lectura.
        for ws in writer.book.worksheets:
            for col_cells in ws.columns:
                max_len = 0
                for cell in col_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(value))
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)

    return out.getvalue()

def excel_bytes(df, rutas=None):
    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.drop(columns=["_ID_"], errors="ignore").to_excel(writer, sheet_name="SOLICITUDES", index=False)
        if rutas is not None and not rutas.empty:
            rutas.drop(columns=["_ID_RUTA_"], errors="ignore").to_excel(writer, sheet_name=RUTA_SHEET, index=False)
            formatear_fechas_excel(
                writer.sheets[RUTA_SHEET],
                ["FECHA LIMITE DE INGRESO (SE1)", "FECHA LIMITE DE INGRESO (SR1)", "FECHA DE RECOJO"]
            )

        formatear_fechas_excel(writer.sheets["SOLICITUDES"], ["FECHA DE INGRESO"])
    return out.getvalue()



def excel_bytes_informe_cajas_cintillos(
    df_salidas,
    df_movimientos=None,
    titulo_filtro="",
    fecha_desde=None,
    fecha_hasta=None,
    datos_kardex=None,
):
    """Genera el informe mensual tipo KARDEX para CAJAS RESGUARDO.

    Cambios del informe:
    - Permite completar COD. CENTRO COSTO, PRECIO y TIPO DE CAMBIO
      despues de seleccionar el mes.
    - TOTAL se calcula automaticamente como CANT. SALIDA x PRECIO.
    - Incluye en TRAZABILIDAD los ingresos de CAJAS NUEVAS registrados
      en movimientos_stock.
    - Los CINTILLOS se muestran en RESUMEN, pero no se mezclan con los
      articulos de caja BASE/TAPA de la trazabilidad.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    out = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "TRAZABILIDAD"

    azul = "000099"
    azul_claro = "DDEBF7"
    blanco = "FFFFFF"
    borde = Side(style="thin", color="7F7F7F")
    borde_fila = Side(style="hair", color="D9D9D9")

    encabezados = [
        "Código Artículo", "Artículo", "Cod. Centro Costo",
        "Descripcion de Centro de Costo", "Fecha de Ingreso", "Doc. Ingreso",
        "Fecha Salida", "Doc. Salida", "Transacción", "Cant. Entrada",
        "Cant. Salida", "Moneda", "Precio", "Tipo de Cambio", "Total",
        "Nro. Orden de Compra", "Cod. Proveedor", "Proveedor",
        "Guía de Compra", "Doc. Compra", "Fecha Emisión",
        "Fecha Recepción", "Fecha Vencimiento"
    ]

    # ---------------------------------------------------------
    # RANGO DEL INFORME
    # ---------------------------------------------------------
    if fecha_desde is not None and fecha_hasta is not None:
        desde = pd.to_datetime(fecha_desde).strftime("%d/%m/%Y")
        hasta = pd.to_datetime(fecha_hasta).strftime("%d/%m/%Y")
        desde_dt = pd.to_datetime(fecha_desde)
        hasta_dt = pd.to_datetime(fecha_hasta)
    else:
        fechas = pd.to_datetime(
            df_salidas.get("FECHA DE SALIDA", pd.Series(dtype=object)),
            errors="coerce"
        ).dropna() if df_salidas is not None else pd.Series(dtype="datetime64[ns]")
        if len(fechas):
            desde_dt = fechas.min()
            hasta_dt = fechas.max()
            desde = desde_dt.strftime("%d/%m/%Y")
            hasta = hasta_dt.strftime("%d/%m/%Y")
        else:
            hoy = pd.Timestamp(fecha_local_hoy())
            desde_dt = hasta_dt = hoy
            desde = hasta = hoy.strftime("%d/%m/%Y")

    # ---------------------------------------------------------
    # DATOS INTRODUCIDOS PARA EL KARDEX
    # La llave principal es _ID_CAJA_. Si no existe, se usa una
    # combinacion estable de WO + cliente + agencia + fecha.
    # ---------------------------------------------------------
    detalles = {}
    if datos_kardex is not None and not datos_kardex.empty:
        for _, d in datos_kardex.iterrows():
            key = d.get("_ID_CAJA_")
            if pd.isna(key) or str(key).strip() == "":
                key = "|".join([
                    str(d.get("NRO. WORKORDER", "") or "").strip(),
                    str(d.get("CLIENTE", "") or "").strip(),
                    str(d.get("AGENCIA", "") or "").strip(),
                    str(d.get("FECHA DE SALIDA", "") or "").strip(),
                ])
            detalles[str(key)] = {
                "cod_centro": str(d.get("COD. CENTRO COSTO", "") or "").strip(),
                "precio_base": float(pd.to_numeric(d.get("PRECIO BASE", 0), errors="coerce") or 0),
                "precio_tapa": float(pd.to_numeric(d.get("PRECIO TAPA", 0), errors="coerce") or 0),
                "tipo_cambio": float(pd.to_numeric(d.get("TIPO DE CAMBIO", 1), errors="coerce") or 1),
            }

    def detalle_para_salida(r):
        key = r.get("_ID_CAJA_")
        if pd.isna(key) or str(key).strip() == "":
            key = "|".join([
                str(r.get("NRO. WORKORDER", "") or "").strip(),
                str(r.get("CLIENTE", "") or "").strip(),
                str(r.get("AGENCIA", "") or "").strip(),
                str(r.get("FECHA DE SALIDA", "") or "").strip(),
            ])
        return detalles.get(str(key), {"cod_centro": "", "precio_base": 0.0, "precio_tapa": 0.0, "tipo_cambio": 1.0})

    # ---------------------------------------------------------
    # CABECERA
    # ---------------------------------------------------------
    ws.merge_cells("A1:C1")
    ws["A1"] = "POLYSISTEMAS SRL."
    ws["A1"].font = Font(name="Arial", size=11, bold=True)
    ws["A1"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A3:W3")
    ws["A3"] = "Trazabilidad de Movimientos"
    ws["A3"].font = Font(name="Arial", size=11, bold=True, color=blanco)
    ws["A3"].fill = PatternFill("solid", fgColor=azul)
    ws["A3"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A4:W4")
    rango_txt = f"Rango de Fechas (Ingresos \\ Salidas).  Desde: {desde} Hasta: {hasta}"
    if titulo_filtro:
        rango_txt += f"  |  {titulo_filtro}"
    ws["A4"] = rango_txt
    ws["A4"].font = Font(name="Arial", size=10, bold=True, color=blanco)
    ws["A4"].fill = PatternFill("solid", fgColor=azul)
    ws["A4"].alignment = Alignment(horizontal="left")
    ws["A4"].border = Border(bottom=borde)

    ws.merge_cells("C6:O6")
    ws["C6"] = "Movimiento de Almacén"
    ws["C6"].font = Font(name="Arial", size=10, bold=True, color=blanco)
    ws["C6"].fill = PatternFill("solid", fgColor=azul)
    ws["C6"].alignment = Alignment(horizontal="center")

    for col, encabezado in enumerate(encabezados, 1):
        c = ws.cell(7, col, encabezado)
        c.font = Font(name="Arial", size=9, bold=True, color=blanco)
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=borde)

    # ---------------------------------------------------------
    # TRAZABILIDAD
    # 1) INGRESOS DE CAJAS NUEVAS desde movimientos_stock
    # 2) SALIDAS: cada caja = BASE + TAPA
    # ---------------------------------------------------------
    filas = []

    if df_movimientos is not None and not df_movimientos.empty:
        mov = df_movimientos.copy()
        fechas_mov = pd.to_datetime(mov.get("fecha"), errors="coerce")
        tipo_mov = mov.get("tipo_movimiento", "").astype(str).str.upper()
        cajas_nuevas = pd.to_numeric(mov.get("cajas_nuevas", 0), errors="coerce").fillna(0)
        mask_ingreso = (
            fechas_mov.notna()
            & (fechas_mov >= desde_dt)
            & (fechas_mov <= hasta_dt)
            & tipo_mov.eq("INGRESO")
            & (cajas_nuevas > 0)
        )

        for _, m in mov.loc[mask_ingreso].iterrows():
            fecha = convertir_fecha(m.get("fecha"))
            cantidad = entero_seguro(m.get("cajas_nuevas"))
            proveedor = str(m.get("proveedor", "") or "").strip()
            observaciones = str(m.get("observaciones", "") or "").strip()
            doc_ingreso = "ING-STOCK"
            if observaciones:
                doc_ingreso = f"ING-STOCK | {observaciones}"

            for codigo, articulo in [
                ("23001001", "BASE P. CAJA RESGUARDO 1.2P - IMPRESAS - LARGO 38.2 X ANCHO 30 X 25.3 CMS"),
                ("23001002", "TAPA P. CAJA RESGUARDO 1.2P - IMPRESAS - LARGO 39.5 X ANCHO 30.8 X 7 CMS"),
            ]:
                filas.append({
                    "codigo": codigo,
                    "articulo": articulo,
                    "cod_centro": "",
                    "centro": "INGRESO DE CAJAS NUEVAS",
                    "fecha_ingreso": fecha,
                    "doc_ingreso": doc_ingreso,
                    "fecha_salida": None,
                    "doc_salida": "",
                    "transaccion": "INGRESO",
                    "cant_entrada": cantidad,
                    "cant_salida": None,
                    "moneda": "",
                    "precio": None,
                    "tipo_cambio": None,
                    "total": None,
                    "proveedor": proveedor,
                })

    if df_salidas is not None and not df_salidas.empty:
        for _, r in df_salidas.iterrows():
            fecha = convertir_fecha(r.get("FECHA DE SALIDA"))
            wo = str(r.get("NRO. WORKORDER", "") or "").strip()
            cliente = str(r.get("CLIENTE", "") or "").strip()
            agencia = str(r.get("AGENCIA", "") or "").strip()
            centro = f"{cliente} - {agencia}" if agencia else cliente
            cajas = entero_seguro(r.get("CANTIDAD DE CAJAS SOLICITADAS"))
            detalle = detalle_para_salida(r)

            if cajas <= 0:
                continue

            for codigo, articulo, precio in [
                ("23001001", "BASE P. CAJA RESGUARDO 1.2P - IMPRESAS - LARGO 38.2 X ANCHO 30 X 25.3 CMS", detalle["precio_base"]),
                ("23001002", "TAPA P. CAJA RESGUARDO 1.2P - IMPRESAS - LARGO 39.5 X ANCHO 30.8 X 7 CMS", detalle["precio_tapa"]),
            ]:
                filas.append({
                    "codigo": codigo,
                    "articulo": articulo,
                    "cod_centro": detalle["cod_centro"],
                    "centro": centro,
                    "fecha_ingreso": None,
                    "doc_ingreso": "",
                    "fecha_salida": fecha,
                    "doc_salida": wo,
                    "transaccion": "AF",
                    "cant_entrada": None,
                    "cant_salida": cajas,
                    "moneda": "BOB",
                    "precio": precio,
                    "tipo_cambio": detalle["tipo_cambio"],
                    "total": round(cajas * precio, 2),
                    "proveedor": "",
                })

    for fila_num, r in enumerate(filas, 8):
        valores = [
            r["codigo"], r["articulo"], r["cod_centro"], r["centro"],
            r["fecha_ingreso"], r["doc_ingreso"], r["fecha_salida"], r["doc_salida"],
            r["transaccion"], r["cant_entrada"], r["cant_salida"], r["moneda"],
            r["precio"], r["tipo_cambio"], r["total"],
            "", "", r["proveedor"], "", "", None, None, None,
        ]
        for col, valor in enumerate(valores, 1):
            c = ws.cell(fila_num, col, valor)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=(col in (2, 4)))
            c.border = Border(bottom=borde_fila)

        for col in (5, 7, 21, 22, 23):
            ws.cell(fila_num, col).number_format = "DD/MM/YYYY"
        for col in (13, 14, 15):
            ws.cell(fila_num, col).number_format = '#,##0.00'

    if not filas:
        ws.cell(8, 1, "SIN REGISTROS DE MOVIMIENTOS")
        ws.merge_cells("A8:W8")
        ws["A8"].alignment = Alignment(horizontal="center")
        ws["A8"].font = Font(name="Arial", size=10, italic=True)

    anchos = [14, 62, 15, 36, 14, 18, 14, 16, 13, 13, 13, 12, 12, 13, 14, 18, 15, 22, 16, 16, 14, 14, 16]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[7].height = 42
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:W{max(7, 7 + len(filas))}"
    ws.sheet_view.showGridLines = False

    # ---------------------------------------------------------
    # RESUMEN
    # ---------------------------------------------------------
    resumen = wb.create_sheet("RESUMEN")
    resumen.sheet_view.showGridLines = False
    resumen.merge_cells("A1:D1")
    resumen["A1"] = "INFORME DE CONTROL DE CAJAS Y CINTILLOS"
    resumen["A1"].font = Font(name="Arial", size=14, bold=True, color=blanco)
    resumen["A1"].fill = PatternFill("solid", fgColor=azul)
    resumen["A1"].alignment = Alignment(horizontal="center")

    total_cajas = (
        int(pd.to_numeric(df_salidas.get("CANTIDAD DE CAJAS SOLICITADAS"), errors="coerce").fillna(0).sum())
        if df_salidas is not None and not df_salidas.empty and "CANTIDAD DE CAJAS SOLICITADAS" in df_salidas.columns else 0
    )
    total_cintillos = (
        int(pd.to_numeric(df_salidas.get("CANTIDAD DE CINTILLOS"), errors="coerce").fillna(0).sum())
        if df_salidas is not None and not df_salidas.empty and "CANTIDAD DE CINTILLOS" in df_salidas.columns else 0
    )
    total_cajas_nuevas = 0
    if df_movimientos is not None and not df_movimientos.empty:
        fechas_mov = pd.to_datetime(df_movimientos.get("fecha"), errors="coerce")
        tipo_mov = df_movimientos.get("tipo_movimiento", "").astype(str).str.upper()
        nuevas = pd.to_numeric(df_movimientos.get("cajas_nuevas", 0), errors="coerce").fillna(0)
        total_cajas_nuevas = int(nuevas[
            fechas_mov.notna() & (fechas_mov >= desde_dt) & (fechas_mov <= hasta_dt) & tipo_mov.eq("INGRESO")
        ].sum())

    total_registros = len(df_salidas) if df_salidas is not None else 0
    total_movimientos_kardex = len(filas)
    total_valorizado = round(sum(float(r.get("total") or 0) for r in filas if r.get("transaccion") == "AF"), 2)

    resumen_data = [
        ("Rango del informe", f"{desde} al {hasta}"),
        ("Registros de salida", total_registros),
        ("Total de cajas solicitadas", total_cajas),
        ("Cajas nuevas ingresadas", total_cajas_nuevas),
        ("Total de cintillos", total_cintillos),
        ("Movimientos en trazabilidad", total_movimientos_kardex),
        ("Valor total de salidas", total_valorizado),
        ("Artículos por caja", "BASE 23001001 + TAPA 23001002"),
        ("Cintillos", "Incluidos en el RESUMEN; no se mezclan con BASE/TAPA"),
        ("Generado", datetime.now(ZONA_HORARIA_APP).strftime("%d/%m/%Y %H:%M")),
    ]
    for rr, (ind, val) in enumerate(resumen_data, 3):
        resumen.cell(rr, 1, ind)
        resumen.cell(rr, 2, val)
        for cc in (1, 2):
            resumen.cell(rr, cc).font = Font(name="Arial", size=10, bold=(cc == 1))
            resumen.cell(rr, cc).fill = PatternFill("solid", fgColor=azul_claro if rr % 2 else "FFFFFF")
            resumen.cell(rr, cc).border = Border(bottom=borde)
            resumen.cell(rr, cc).alignment = Alignment(vertical="center", wrap_text=True)
    resumen["B9"].number_format = '#,##0.00'

    resumen.column_dimensions["A"].width = 34
    resumen.column_dimensions["B"].width = 55
    resumen.column_dimensions["C"].width = 4
    resumen.column_dimensions["D"].width = 4

    wb.save(out)
    return out.getvalue()

def excel_bytes_historial_salidas(df):
    """Genera un archivo Excel del historial de salidas de cajas y cintillos."""
    out = BytesIO()
    datos = df.drop(columns=["_ID_CAJA_"], errors="ignore").copy()

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        datos.to_excel(writer, sheet_name="HISTORIAL SALIDAS", index=False)
        formatear_fechas_excel(writer.sheets["HISTORIAL SALIDAS"], ["FECHA DE SALIDA"])
        ws = writer.sheets["HISTORIAL SALIDAS"]

        # Ajustar el ancho de las columnas para facilitar la lectura.
        for col_cells in ws.columns:
            max_len = 0
            for cell in col_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[col_cells[0].column_letter].width = min(
                max(max_len + 2, 12), 45
            )

    return out.getvalue()

# =========================================================
# LOGIN
# =========================================================
def obtener_usuarios_login():
    try:
        bloque = st.secrets.get("usuarios", {})
        if hasattr(bloque, "items"):
            return {str(k).strip().lower(): str(v) for k, v in bloque.items()}
    except Exception:
        pass
    return {}


def nombres_usuarios():
    return {
        "admin": "admin",
        "alfredo": "Alfredo",
        "jonathan": "Jonathan",
        "luis": "Luis",
        "invitado": "Invitado",
    }


def obtener_rol():
    usuario = str(st.session_state.get("usuario_login", "") or "").strip().lower()
    return {
        "admin": "admin",
        "alfredo": "usuario",
        "jonathan": "usuario",
        "luis": "usuario",
        "invitado": "invitado",
    }.get(usuario, "usuario")


def es_admin():
    return obtener_rol() == "admin"


def es_invitado():
    return obtener_rol() == "invitado"


def puede_modificar():
    # ADMIN y usuarios normales pueden registrar/editar.
    # INVITADO solamente puede consultar.
    return obtener_rol() in ("admin", "usuario")


def mostrar_login():
    st.markdown(
        '<div class="login-wrapper"><div class="login-icon">🔐</div><div class="login-title">Iniciar sesión</div><div class="login-subtitle">Acceso al sistema de despacho</div></div>',
        unsafe_allow_html=True,
    )
    usuarios = obtener_usuarios_login()
    if not usuarios:
        st.error("❌ No hay usuarios configurados en Streamlit Secrets.")
        return
    _, centro, _ = st.columns([1, 1.25, 1])
    with centro:
        with st.form("login_jc_control"):
            usuario = st.text_input("👤 Usuario", placeholder="Ingrese su usuario")
            password = st.text_input("🔑 Contraseña", type="password", placeholder="Ingrese su contraseña")
            ingresar = st.form_submit_button("🔐 INGRESAR", use_container_width=True)
    if ingresar:
        key = usuario.strip().lower()
        if key in usuarios and password == usuarios[key]:
            st.session_state.autenticado = True
            st.session_state.usuario_login = key
            st.session_state.usuario_nombre = nombres_usuarios().get(key, usuario.strip())
            st.rerun()
        else:
            st.error("❌ Usuario o contraseña incorrectos.")

# =========================================================
# FORMULARIO PROGRAMACION DE RUTAS
# =========================================================
def formulario_nueva_ruta():
    st.markdown("### ➕ Registrar programación de ruta")
    with st.form(f"ruta_form_{st.session_state.ruta_form_version}", clear_on_submit=False):
        a, b = st.columns(2)
        with a:
            regional = st.selectbox("🌎 Regional", REGIONALES)
            centro_acopio = st.text_input("🏢 Centro de acopio", placeholder="Ej.: POLYSISTEMAS")
            agencias = st.text_input("🏪 Agencias", placeholder="Ej.: San Miguel, Zona Sur")
        with b:
            fecha_se1 = st.date_input("📅 Fecha límite de ingreso (SE1)", value=fecha_local_hoy(), format="DD/MM/YYYY")
            fecha_sr1 = st.date_input("📅 Fecha límite de ingreso (SR1)", value=fecha_local_hoy(), format="DD/MM/YYYY")
            fecha_recojo = st.date_input("🚚 Fecha de recojo", value=fecha_local_hoy(), format="DD/MM/YYYY")

        x, y = st.columns([4, 1])
        with x:
            guardar = st.form_submit_button("💾 GUARDAR PROGRAMACIÓN", use_container_width=True)
        with y:
            limpiar = st.form_submit_button("✕ CANCELAR", type="primary", use_container_width=True)

    if limpiar:
        liberar_bloqueo_edicion()
        st.session_state.ruta_form_version += 1
        st.rerun()

    if guardar:
        if regional == REGIONALES[0]:
            st.error("Selecciona una Regional antes de guardar.")
            return
        if not centro_acopio.strip():
            st.error("Ingresa el Centro de Acopio antes de guardar.")
            return
        if not agencias.strip():
            st.error("Ingresa las Agencias antes de guardar.")
            return

        if fecha_se1 > fecha_recojo or fecha_sr1 > fecha_recojo:
            st.error("La fecha de recojo no puede ser anterior a las fechas límite de ingreso.")
            return

        try:
            datos = {
                "REGIONAL": regional,
                "CENTRO DE ACOPIO": centro_acopio.strip(),
                "AGENCIAS": agencias.strip(),
                "FECHA LIMITE DE INGRESO (SE1)": fecha_se1.isoformat(),
                "FECHA LIMITE DE INGRESO (SR1)": fecha_sr1.isoformat(),
                "FECHA DE RECOJO": fecha_recojo.isoformat(),
            }
            guardar_ruta_supabase(datos)
            st.session_state.rutas = cargar_rutas_supabase()
            liberar_bloqueo_edicion()
            st.session_state.ruta_form_version += 1
            st.success("✅ Programación de ruta guardada correctamente en Supabase.")
            st.rerun()
        except Exception as e:
            liberar_bloqueo_edicion()
            st.error(f"❌ No se pudo guardar la programación: {e}")


# =========================================================
# CABECERA Y LOGIN
# =========================================================
st.markdown('<div class="titulo">📝 JC CONTROL DE SOLICITUDES — DESPACHO</div>', unsafe_allow_html=True)

if not st.session_state.autenticado:
    mostrar_login()
    st.stop()

# =========================================================
# CIERRE AUTOMÁTICO POR INACTIVIDAD
# 8 minutos sin interacción del usuario.
# =========================================================
_refresh_count = st_autorefresh(interval=AUTO_REFRESH_INTERVAL_MS, key="control_inactividad")
_ultimo_refresh = st.session_state.get("_ultimo_refresh_inactividad")
_ahora = time.time()
if "_ultima_actividad" not in st.session_state:
    st.session_state._ultima_actividad = _ahora
elif _ultimo_refresh is not None and _refresh_count == _ultimo_refresh:
    st.session_state._ultima_actividad = _ahora
st.session_state._ultimo_refresh_inactividad = _refresh_count

_inactividad = _ahora - st.session_state._ultima_actividad
if _inactividad >= INACTIVITY_TIMEOUT_SECONDS:
    liberar_bloqueo_edicion()
    for k in ["rows", "rutas", "cajas", "movimientos_stock", "comisiones"]:
        st.session_state[k] = None
    st.session_state.editing = None
    st.session_state.editing_key = ""
    st.session_state.usuario_nombre = ""
    st.session_state.autenticado = False
    st.session_state.page = "solicitudes"
    st.session_state.pop("_ultima_actividad", None)
    st.session_state.pop("_ultimo_refresh_inactividad", None)
    st.rerun()

if _inactividad >= 5 * 60:
    minutos_restantes = max(1, int((INACTIVITY_TIMEOUT_SECONDS - _inactividad + 59) // 60))
    st.warning(f"⏳ Tu sesión se cerrará por inactividad en aproximadamente {minutos_restantes} minuto(s).")

u1, u2, u3 = st.columns([1, 3, 1])
with u1:
    st.markdown(f"**👤 Usuario:** `{st.session_state.usuario_nombre}`")
    if es_admin():
        st.caption("🔴 ADMIN — ACCESO TOTAL")
    elif es_invitado():
        st.caption("👁️ INVITADO — SOLO CONSULTA")
    else:
        st.caption("🟢 USUARIO — REGISTRAR Y EDITAR")
with u2:
    st.markdown("**🗄️ Base de datos:** `PostgreSQL`")
with u3:
    if st.button("🚪 SALIR", use_container_width=True):
        liberar_bloqueo_edicion()
        for k in ["rows", "rutas", "cajas", "movimientos_stock", "comisiones"]:
            st.session_state[k] = None
        st.session_state.editing = None
        st.session_state.editing_key = ""
        st.session_state.usuario_nombre = ""
        st.session_state.autenticado = False
        st.session_state.page = "solicitudes"
        st.rerun()

# =========================================================
# CARGA BAJO DEMANDA
# =========================================================
# Antes se cargaban solicitudes, rutas, cajas, stock y comisiones
# simultáneamente al entrar al sistema. Eso hacía lenta la primera carga
# y también la navegación. Ahora cada módulo consulta solamente sus datos
# cuando el usuario entra a esa sección.
def asegurar_datos_pagina():
    pagina = st.session_state.page
    try:
        if pagina == "solicitudes":
            if st.session_state.rows is None:
                st.session_state.rows = cargar_solicitudes_supabase()
        elif pagina == "rutas":
            if st.session_state.rutas is None:
                st.session_state.rutas = cargar_rutas_supabase()
        elif pagina == "cajas":
            if st.session_state.cajas is None:
                st.session_state.cajas = cargar_cajas_supabase()
            if st.session_state.movimientos_stock is None:
                st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
        elif pagina == "comisiones":
            if st.session_state.comisiones is None:
                st.session_state.comisiones = cargar_comisiones_supabase()
    except Exception as e:
        st.error(f"❌ No se pudieron cargar los datos desde Supabase: {e}")
        st.stop()


# Catálogos: se inicializan con un placeholder local y solamente se consulta
# Supabase para los catálogos necesarios en la página actual.
CAT_CLIENTES = ["Seleccione una opcion"]
CAT_TIPOS = ["Seleccione una opcion"]
CAT_CENTROS_COSTO = ["Seleccione una opcion"]
CAT_ESTADOS = ["Seleccione una opcion"]
CAT_DIRECCIONES = ["Seleccione una opcion"]
CAT_AGENCIAS = ["Seleccione una opcion"]
CAT_TIPOS_SALIDA_COMISION = ["Seleccione una opcion"]
CAT_OPERADORES_COMISION = ["Seleccione una opcion"]

def cargar_catalogos_pagina():
    global CAT_CLIENTES, CAT_TIPOS, CAT_CENTROS_COSTO, CAT_ESTADOS
    global CAT_DIRECCIONES, CAT_AGENCIAS, CAT_TIPOS_SALIDA_COMISION, CAT_OPERADORES_COMISION

    pagina = st.session_state.page
    try:
        if pagina == "solicitudes":
            CAT_CLIENTES = cargar_catalogo("CLIENTE")
            CAT_TIPOS = cargar_catalogo("TIPO_SOLICITUD")
            CAT_CENTROS_COSTO = cargar_catalogo("CENTRO_COSTO")
            CAT_ESTADOS = cargar_catalogo("ESTADO_SOLICITUD")
            CAT_DIRECCIONES = cargar_catalogo("DIRECCION")
        elif pagina == "cajas":
            CAT_CLIENTES = cargar_catalogo("CLIENTE")
            CAT_AGENCIAS = cargar_catalogo("AGENCIA")
        elif pagina == "comisiones":
            CAT_TIPOS_SALIDA_COMISION = cargar_catalogo("TIPO_SALIDA_COMISION")
            CAT_OPERADORES_COMISION = cargar_catalogo("OPERADOR_COMISION")
        elif pagina == "catalogos":
            CAT_CLIENTES = cargar_catalogo("CLIENTE")
            CAT_TIPOS = cargar_catalogo("TIPO_SOLICITUD")
            CAT_CENTROS_COSTO = cargar_catalogo("CENTRO_COSTO")
            CAT_ESTADOS = cargar_catalogo("ESTADO_SOLICITUD")
            CAT_DIRECCIONES = cargar_catalogo("DIRECCION")
            CAT_AGENCIAS = cargar_catalogo("AGENCIA")
            CAT_TIPOS_SALIDA_COMISION = cargar_catalogo("TIPO_SALIDA_COMISION")
            CAT_OPERADORES_COMISION = cargar_catalogo("OPERADOR_COMISION")
    except Exception as e:
        st.error(f"❌ No se pudieron cargar los catálogos desde Supabase: {e}")
        st.info("Verifica que exista la tabla `catalogos_app` y que el usuario de Supabase tenga permisos de SELECT.")
        st.stop()

# =========================================================
# NAVEGACION
# =========================================================
c1, c2, c3, c4 = st.columns([1, 1.35, 1.45, 1.55])
with c1:
    if st.button("🔄 ACTUALIZAR", use_container_width=True):
        try:
            pagina_actual = st.session_state.page
            if pagina_actual == "solicitudes":
                st.session_state.rows = cargar_solicitudes_supabase()
            elif pagina_actual == "rutas":
                st.session_state.rutas = cargar_rutas_supabase()
            elif pagina_actual == "cajas":
                st.session_state.cajas = cargar_cajas_supabase()
                st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
            elif pagina_actual == "comisiones" and es_admin():
                st.session_state.comisiones = cargar_comisiones_supabase()
            elif pagina_actual == "catalogos":
                cargar_catalogo.clear()
            st.session_state.editing = None
            st.session_state.editing_key = ""
            st.toast("🔄 Datos de la sección actualizados desde Supabase.", icon="🔄")
            st.rerun()
        except Exception as e:
            st.error(f"❌ No se pudieron actualizar los datos: {e}")
with c2:
    if st.button("🗓️ PROGRAMACIÓN DE RUTAS", use_container_width=True):
        liberar_bloqueo_edicion()
        st.session_state.editing = None
        st.session_state.editing_key = ""
        st.session_state.page = "rutas"
        st.rerun()
with c3:
    if st.button("📦 CONTROL DE CAJAS Y CINTILLOS", use_container_width=True):
        liberar_bloqueo_edicion()
        st.session_state.editing = None
        st.session_state.editing_key = ""
        st.session_state.page = "cajas"
        st.rerun()

with c4:
    # No mostrar este botón mientras ya estamos dentro de la sección de comisiones,
    # para evitar que aparezca duplicado junto al botón/expander de registro.
    if es_admin() and st.session_state.page != "comisiones":
        if st.button("🚶 CONTROL DE SALIDAS A COMISIÓN", use_container_width=True):
            liberar_bloqueo_edicion()
            st.session_state.editing = None
            st.session_state.editing_key = ""
            st.session_state.page = "comisiones"
            st.rerun()

if es_admin() and st.session_state.page not in ["catalogos"]:
    if st.button("⚙️ ADMINISTRAR CATÁLOGOS", use_container_width=True):
        liberar_bloqueo_edicion()
        st.session_state.page = "catalogos"
        st.rerun()

# =========================================================
# CARGAR SOLO LO NECESARIO PARA LA SECCIÓN ACTUAL
# =========================================================
asegurar_datos_pagina()
cargar_catalogos_pagina()

# =========================================================
# PAGINA ADMINISTRAR CATÁLOGOS
# =========================================================
if st.session_state.page == "catalogos":
    if not es_admin():
        st.error("⛔ Esta sección es exclusiva para el usuario ADMIN.")
        st.stop()

    st.subheader("⚙️ ADMINISTRAR CATÁLOGOS")
    st.caption("Gestiona los valores desde Supabase. Desactivar un valor no modifica los registros históricos.")

    # Organización compacta por módulo: cada pestaña contiene únicamente
    # los catálogos relacionados con esa sección del sistema.
    grupos_catalogos = {
        "📁 SOLICITUDES": [
            ("CLIENTES", "CLIENTE"),
            ("TIPOS DE SOLICITUD", "TIPO_SOLICITUD"),
            ("CENTROS DE COSTO", "CENTRO_COSTO"),
            ("ESTADOS DE SOLICITUD", "ESTADO_SOLICITUD"),
            ("DIRECCIONES", "DIRECCION"),
        ],
        "📦 CAJAS Y CINTILLOS": [
            ("AGENCIAS", "AGENCIA"),
            ("CLIENTES", "CLIENTE"),
        ],
        "🚶 COMISIÓN": [
            ("OPERADORES", "OPERADOR_COMISION"),
            ("TIPOS DE SALIDA", "TIPO_SALIDA_COMISION"),
        ],
    }

    def administrar_catalogo_compacto(nombre_cat, categoria, sufijo):
        valores_activos = cargar_catalogo(categoria, incluir_placeholder=False)
        st.markdown(f"**{nombre_cat}** · {len(valores_activos)} activos")

        a1, a2 = st.columns([3.4, 1])
        with a1:
            nuevo_valor = st.text_input(
                "Nuevo valor",
                placeholder="Escribe el nuevo elemento...",
                key=f"admin_nuevo_{sufijo}",
                label_visibility="collapsed",
            )
        with a2:
            agregar = st.button(
                "➕ AGREGAR",
                use_container_width=True,
                key=f"admin_agregar_{sufijo}",
            )

        if agregar:
            try:
                agregar_catalogo(categoria, nuevo_valor)
                st.success(f"✅ Valor agregado a {nombre_cat}.")
                st.rerun()
            except Exception as e:
                st.error(f"❌ No se pudo agregar: {e}")

        if valores_activos:
            b1, b2 = st.columns([3.4, 1])
            with b1:
                valor_eliminar = st.selectbox(
                    "Valor a desactivar",
                    valores_activos,
                    key=f"admin_eliminar_{sufijo}",
                    label_visibility="collapsed",
                )
            with b2:
                desactivar = st.button(
                    "🗑️ DESACTIVAR",
                    type="primary",
                    use_container_width=True,
                    key=f"admin_desactivar_{sufijo}",
                )
            if desactivar:
                try:
                    desactivar_catalogo(categoria, valor_eliminar)
                    st.success("✅ Valor desactivado. Los registros históricos se mantienen intactos.")
                    st.rerun()
                except Exception as e:
                    st.error(f"❌ No se pudo desactivar: {e}")

            with st.expander(f"📋 Ver valores activos ({len(valores_activos)})", expanded=False):
                st.dataframe(
                    pd.DataFrame({"VALOR": valores_activos}),
                    use_container_width=True,
                    hide_index=True,
                    height=220,
                )
        else:
            st.info("No hay valores activos en este catálogo.")

    tabs = st.tabs(list(grupos_catalogos.keys()))
    for tab, (grupo_nombre, lista_catalogos) in zip(tabs, grupos_catalogos.items()):
        with tab:
            for idx, (nombre_cat, categoria) in enumerate(lista_catalogos):
                administrar_catalogo_compacto(nombre_cat, categoria, f"{grupo_nombre}_{idx}")
                if idx < len(lista_catalogos) - 1:
                    st.divider()

    st.divider()
    if st.button("⬅ VOLVER A SOLICITUDES", key="volver_catalogos", use_container_width=True):
        st.session_state.page = "solicitudes"
        st.rerun()

    st.stop()

# =========================================================
# PAGINA CONTROL DE SALIDAS A COMISIÓN — SOLO ADMIN
# =========================================================
if st.session_state.page == "comisiones":
    if not es_admin():
        st.error("⛔ Esta sección es exclusiva para el usuario ADMIN.")
        st.stop()

    st.subheader("🚶 CONTROL DE SALIDAS A COMISIÓN")
    st.caption("Registro y control del personal que sale en comisión, por ejemplo para entrega o recojo de documentos.")

    if st.button("⬅ VOLVER A SOLICITUDES", key="volver_comisiones"):
        st.session_state.page = "solicitudes"
        st.rerun()

    # =====================================================
    # REGISTRAR NUEVA SALIDA: EXPANDER CERRADO POR DEFECTO
    # =====================================================
    titulo_registro = (
        "➕ REGISTRAR SALIDA A COMISIÓN"
        if st.session_state.comision_form_version % 2 == 0
        else "➕ REGISTRAR SALIDA A COMISIÓN\u2060"
    )
    with st.expander(
        titulo_registro,
        expanded=st.session_state.get("mostrar_form_comision", False),
    ):
        # Estos checkboxes están FUERA del st.form para que al seleccionarlos
        # la pantalla se actualice inmediatamente y aparezca la hora actual.
        a0, b0 = st.columns(2)
        with a0:
            registrar_hora_ingreso = st.checkbox(
                "Registrar HORA DE INGRESO",
                value=False,
                key=f"registrar_hora_ingreso_{st.session_state.comision_form_version}",
            )
        with b0:
            registrar_hora_salida = st.checkbox(
                "Registrar HORA DE SALIDA",
                value=False,
                key=f"registrar_hora_salida_{st.session_state.comision_form_version}",
            )

        hora_ingreso = None
        hora_salida = None
        if registrar_hora_ingreso or registrar_hora_salida:
            h1, h2 = st.columns(2)
            with h1:
                if registrar_hora_ingreso:
                    hora_ingreso = st.time_input(
                        "🕘 HORA DE INGRESO",
                        value=hora_actual_local(),
                        step=60,
                        key=f"hora_ingreso_{st.session_state.comision_form_version}",
                    )
                    st.caption("Hora actual cargada automáticamente. Puedes modificarla.")
            with h2:
                if registrar_hora_salida:
                    hora_salida = st.time_input(
                        "🕘 HORA DE SALIDA",
                        value=hora_actual_local(),
                        step=60,
                        key=f"hora_salida_{st.session_state.comision_form_version}",
                    )
                    st.caption("Hora actual cargada automáticamente. Puedes modificarla.")

        with st.form(
            f"form_comision_{st.session_state.comision_form_version}",
            clear_on_submit=False,
        ):
            a, b, c = st.columns(3)

            with a:
                operador = st.selectbox("👤 OPERADOR", CAT_OPERADORES_COMISION)
                fecha_comision = st.date_input(
                    "📅 FECHA",
                    value=fecha_local_hoy(),
                    format="DD/MM/YYYY",
                )

            with b:
                tipo_salida = st.selectbox("🚶 TIPO DE SALIDA", CAT_TIPOS_SALIDA_COMISION)

            with c:
                if not (registrar_hora_ingreso or registrar_hora_salida):
                    st.caption("🕘 Activa una de las opciones de hora para registrar el horario.")
                elif registrar_hora_ingreso and registrar_hora_salida:
                    st.caption("🕘 Ingreso y salida registradas.")
                elif registrar_hora_ingreso:
                    st.caption("🕘 Solo se registrará la hora de ingreso.")
                else:
                    st.caption("🕘 Solo se registrará la hora de salida.")

            descripcion = st.text_area(
                "📝 DESCRIPCIÓN",
                placeholder="Ej.: Entrega de documentos en agencia...",
            )

            # Tiempo adicional para compensación. No modifica las horas de ingreso/salida.
            tiempo_adicional_label = st.selectbox(
                "➕ TIEMPO ADICIONAL PARA COMPENSACIÓN",
                list(TIEMPOS_ADICIONALES.keys()),
                key=f"tiempo_adicional_{st.session_state.comision_form_version}",
                help="Permite sumar 30 min, 1 h, 1 h 30 min o 2 h sin cambiar la hora de ingreso ni la hora de salida.",
            )
            tiempo_adicional = TIEMPOS_ADICIONALES[tiempo_adicional_label]

            horas_base = calcular_total_horas(
                fecha_comision, hora_ingreso, hora_salida
            )
            # El adicional solo se aplica cuando existen ingreso y salida.
            tiempo_adicional_aplicado = tiempo_adicional if horas_base is not None else 0.0
            total_horas = (horas_base + tiempo_adicional_aplicado) if horas_base is not None else None
            if total_horas is not None:
                st.info(
                    f"⏱️ TOTAL HORAS: {formatear_horas_minutos(total_horas)} "
                    f"(base: {formatear_horas_minutos(horas_base)} + adicional: {formatear_horas_minutos(tiempo_adicional_aplicado)})"
                )
            else:
                st.caption(
                    "⏱️ TOTAL HORAS: se calculará automáticamente solo cuando registres ambas horas."
                )

            # Un registro nuevo siempre inicia como pendiente de compensación.
            st.caption("📌 Las horas nuevas quedan como PENDIENTES hasta que ADMIN las marque como compensadas.")

            g, x = st.columns([4, 1])
            with g:
                guardar = st.form_submit_button(
                    "💾 GUARDAR SALIDA",
                    use_container_width=True,
                )
            with x:
                cancelar = st.form_submit_button(
                    "✕ CANCELAR",
                    type="primary",
                    use_container_width=True,
                )

        if cancelar:
            st.session_state.mostrar_form_comision = False
            st.session_state.comision_form_version += 1
            st.rerun()

        if guardar:
            if operador == CAT_OPERADORES_COMISION[0]:
                st.error("❌ Selecciona el OPERADOR.")
            elif tipo_salida == CAT_TIPOS_SALIDA_COMISION[0]:
                st.error("❌ Selecciona el TIPO DE SALIDA.")
            elif not descripcion.strip():
                st.error("❌ Ingresa una DESCRIPCIÓN de la comisión.")
            else:
                try:
                    datos = {
                        "OPERADOR": operador,
                        "FECHA": fecha_comision.isoformat(),
                        "TIPO DE SALIDA": tipo_salida,
                        "HORA DE INGRESO": hora_ingreso.strftime("%H:%M:%S") if hora_ingreso is not None else None,
                        "HORA DE SALIDA": hora_salida.strftime("%H:%M:%S") if hora_salida is not None else None,
                        "TIEMPO ADICIONAL": tiempo_adicional_aplicado,
                        "TOTAL HORAS": total_horas,
                        "COMPENSADO": False,
                        "FECHA COMPENSACIÓN": None,
                        "DESCRIPCIÓN": descripcion.strip(),
                    }
                    guardar_comision_supabase(datos)
                    st.session_state.comisiones = cargar_comisiones_supabase()
                    st.session_state.mostrar_form_comision = False
                    st.session_state.comision_form_version += 1
                    st.success("✅ Salida a comisión registrada correctamente en Supabase.")
                    st.rerun()
                except Exception as e:
                    st.error(f"❌ No se pudo guardar la salida a comisión: {e}")

    st.markdown("### 📋 Historial de salidas a comisión")
    try:
        df_com = cargar_comisiones_supabase()
        st.session_state.comisiones = df_com

        if df_com.empty:
            st.info("Aún no existen salidas a comisión registradas.")
        else:
            # =================================================
            # FILTRO ÚNICO DEL HISTORIAL
            # =================================================
            st.markdown("### 🔎 FILTRAR HISTORIAL")
            f1, f2 = st.columns([1, 2])

            with f1:
                filtro_tipo = st.selectbox(
                    "Filtrar por",
                    ["Sin filtro", "OPERADOR", "FECHA", "TIPO DE SALIDA", "DESCRIPCIÓN"],
                    key="filtro_comisiones_tipo",
                )

            df_filtrado = df_com.copy()

            with f2:
                if filtro_tipo == "OPERADOR":
                    operadores_filtro = [
                        x for x in CAT_OPERADORES_COMISION
                        if x != CAT_OPERADORES_COMISION[0]
                    ]
                    valor_filtro = st.selectbox(
                        "Selecciona el operador",
                        ["Todos"] + operadores_filtro,
                        key="filtro_comisiones_operador",
                    )
                    if valor_filtro != "Todos":
                        df_filtrado = df_filtrado[
                            df_filtrado["OPERADOR"].astype(str) == valor_filtro
                        ]

                elif filtro_tipo == "FECHA":
                    st.markdown("**Selecciona el rango de fechas**")
                    rf1, rf2 = st.columns(2)

                    with rf1:
                        fecha_inicio_filtro = st.date_input(
                            "📅 Fecha desde",
                            value=fecha_local_hoy(),
                            format="DD/MM/YYYY",
                            key="filtro_comisiones_fecha_inicio",
                        )

                    with rf2:
                        fecha_fin_filtro = st.date_input(
                            "📅 Fecha hasta",
                            value=fecha_local_hoy(),
                            format="DD/MM/YYYY",
                            key="filtro_comisiones_fecha_fin",
                        )

                    if fecha_inicio_filtro > fecha_fin_filtro:
                        st.warning("⚠️ La fecha inicial no puede ser posterior a la fecha final.")
                        df_filtrado = df_filtrado.iloc[0:0]
                    else:
                        fechas = df_filtrado["FECHA"].apply(convertir_fecha)
                        df_filtrado = df_filtrado[
                            fechas.apply(
                                lambda x: (
                                    x is not None
                                    and fecha_inicio_filtro <= x <= fecha_fin_filtro
                                )
                            )
                        ]

                elif filtro_tipo == "TIPO DE SALIDA":
                    tipos_filtro = [
                        x for x in CAT_TIPOS_SALIDA_COMISION
                        if x != CAT_TIPOS_SALIDA_COMISION[0]
                    ]
                    valor_filtro = st.selectbox(
                        "Selecciona el tipo de salida",
                        ["Todos"] + tipos_filtro,
                        key="filtro_comisiones_tipo_valor",
                    )
                    if valor_filtro != "Todos":
                        df_filtrado = df_filtrado[
                            df_filtrado["TIPO DE SALIDA"].astype(str) == valor_filtro
                        ]

                elif filtro_tipo == "DESCRIPCIÓN":
                    valor_filtro = st.text_input(
                        "Buscar en la descripción",
                        placeholder="Escribe una palabra o frase...",
                        key="filtro_comisiones_descripcion",
                    )
                    if valor_filtro.strip():
                        df_filtrado = df_filtrado[
                            df_filtrado["DESCRIPCIÓN"].astype(str).str.contains(
                                valor_filtro.strip(),
                                case=False,
                                na=False,
                            )
                        ]

            # =================================================
            # RESUMEN DE HORAS DEL MISMO FILTRO
            # =================================================
            horas_serie = pd.to_numeric(
                df_filtrado["TOTAL HORAS"], errors="coerce"
            ).fillna(0)

            compensado_bool = df_filtrado["COMPENSADO"].apply(es_valor_compensado)

            # TOTAL HORAS se guarda como horas decimales en Supabase
            # (por ejemplo, 1.50 = 1 h 30 min). Para el acumulado se
            # convierten a minutos para evitar interpretar 0.50 como 50 minutos.
            # Para el resumen, las horas adicionales solo cuentan cuando
            # el registro tiene hora de ingreso Y hora de salida.
            tiene_horario = (
                df_filtrado["HORA DE INGRESO"].astype(str).str.strip().ne("")
                & df_filtrado["HORA DE SALIDA"].astype(str).str.strip().ne("")
            )
            horas_serie = horas_serie.where(tiene_horario, 0)
            adicional_serie = pd.to_numeric(
                df_filtrado["TIEMPO ADICIONAL"], errors="coerce"
            ).fillna(0).where(tiene_horario, 0)

            minutos_acumulados = int(round(float(horas_serie.sum()) * 60))
            minutos_adicionales = int(round(float(adicional_serie.sum()) * 60))
            minutos_compensados = int(round(float(horas_serie[compensado_bool].sum()) * 60))
            minutos_pendientes = max(0, minutos_acumulados - minutos_compensados)

            horas_acumuladas = minutos_acumulados / 60
            horas_adicionales = minutos_adicionales / 60
            horas_compensadas = minutos_compensados / 60
            horas_pendientes = minutos_pendientes / 60

            st.markdown("### ⏱️ RESUMEN DE HORAS PARA COMPENSACIÓN")
            m1, m2, m3, m4, m5 = st.columns(5)
            m1.metric("📊 REGISTROS FILTRADOS", len(df_filtrado))
            m2.metric("⏱️ HORAS ACUMULADAS", formatear_horas_minutos(horas_acumuladas))
            m3.metric("➕ TIEMPO ADICIONAL", formatear_horas_minutos(horas_adicionales))
            m4.metric("✅ HORAS COMPENSADAS", formatear_horas_minutos(horas_compensadas))
            m5.metric("🕐 HORAS PENDIENTES", formatear_horas_minutos(horas_pendientes))

            if filtro_tipo == "FECHA":
                st.caption(
                    f"📊 Resumen y descarga del rango "
                    f"{fecha_inicio_filtro.strftime('%d/%m/%Y')} al "
                    f"{fecha_fin_filtro.strftime('%d/%m/%Y')}."
                )
            elif filtro_tipo != "Sin filtro":
                st.caption(f"📊 El resumen y la descarga corresponden al filtro seleccionado: {filtro_tipo}.")

            # Resumen por operador para saber el saldo de cada persona.
            resumen_operadores = df_filtrado.copy()
            resumen_operadores["_HORAS_"] = pd.to_numeric(
                resumen_operadores["TOTAL HORAS"], errors="coerce"
            ).fillna(0)
            resumen_operadores["_COMP_"] = compensado_bool.values
            resumen_operadores["_HORAS_COMP_"] = resumen_operadores["_HORAS_"].where(
                resumen_operadores["_COMP_"], 0
            )
            resumen_operadores["_HORAS_PEND_"] = resumen_operadores["_HORAS_"].where(
                ~resumen_operadores["_COMP_"], 0
            )
            resumen_operadores = (
                resumen_operadores.groupby("OPERADOR", dropna=False)[
                    ["_HORAS_", "_HORAS_COMP_", "_HORAS_PEND_"]
                ]
                .sum()
                .reset_index()
                .rename(columns={
                    "_HORAS_": "HORAS ACUMULADAS",
                    "_HORAS_COMP_": "HORAS COMPENSADAS",
                    "_HORAS_PEND_": "HORAS PENDIENTES",
                })
            )
            for col in ["HORAS ACUMULADAS", "HORAS COMPENSADAS", "HORAS PENDIENTES"]:
                resumen_operadores[col] = resumen_operadores[col].apply(formatear_horas_minutos)
            if not resumen_operadores.empty:
                with st.expander("👥 VER SALDO DE HORAS POR OPERADOR", expanded=False):
                    st.dataframe(
                        resumen_operadores,
                        use_container_width=True,
                        hide_index=True,
                    )

            vista_com = df_filtrado.copy()
            vista_com["FECHA"] = vista_com["FECHA"].apply(
                lambda x: convertir_fecha(x).strftime("%d/%m/%Y")
                if convertir_fecha(x) else ""
            )
            vista_com["HORA DE INGRESO"] = (
                vista_com["HORA DE INGRESO"].astype(str)
                .replace("None", "")
                .str[:5]
            )
            vista_com["HORA DE SALIDA"] = (
                vista_com["HORA DE SALIDA"].astype(str)
                .replace("None", "")
                .str[:5]
            )
            adicional_numerico = pd.to_numeric(
                df_filtrado["TIEMPO ADICIONAL"], errors="coerce"
            ).fillna(0)
            vista_com["TIEMPO ADICIONAL"] = [
                etiqueta_tiempo_adicional(adicional)
                if tiene_horario.iloc[i] else ""
                for i, adicional in enumerate(adicional_numerico)
            ]
            vista_com["TOTAL HORAS"] = pd.to_numeric(
                vista_com["TOTAL HORAS"], errors="coerce"
            ).map(lambda x: formatear_horas_minutos(x) if pd.notna(x) else "")
            # COMPENSADO solo se muestra cuando el registro tiene TOTAL HORAS.
            # Si todavía no hay horas calculadas, la celda queda en blanco.
            total_horas_numerico = pd.to_numeric(df_filtrado["TOTAL HORAS"], errors="coerce")
            vista_com["COMPENSADO"] = [
                "" if pd.isna(total)
                else ("✅ SÍ" if es_valor_compensado(comp) else "⏳ NO")
                for total, comp in zip(total_horas_numerico, vista_com["COMPENSADO"])
            ]
            vista_com["FECHA COMPENSACIÓN"] = vista_com["FECHA COMPENSACIÓN"].apply(
                lambda x: convertir_fecha(x).strftime("%d/%m/%Y")
                if convertir_fecha(x) else ""
            )

            # Selección para editar uno o varios registros.
            tabla_com = vista_com.copy()
            tabla_com["☑️ SELECCIONAR"] = False
            tabla_com["✏️ EDITAR"] = False
            tabla_com["🗑️ ELIMINAR"] = False

            resultado_com = st.data_editor(
                tabla_com,
                use_container_width=True,
                hide_index=True,
                height=420,
                key=f"tabla_historial_comisiones_{st.session_state.comision_tabla_version}",
                disabled=[
                    c for c in tabla_com.columns
                    if c not in ["☑️ SELECCIONAR", "✏️ EDITAR", "🗑️ ELIMINAR"]
                ],
                column_config={
                    "_ID_COMISION_": None,
                    "FECHA": st.column_config.TextColumn("FECHA", width=105),
                    "OPERADOR": st.column_config.TextColumn("OPERADOR", width=150),
                    "TIPO DE SALIDA": st.column_config.TextColumn("TIPO DE SALIDA", width=220),
                    "HORA DE INGRESO": st.column_config.TextColumn("HORA INGRESO", width=110),
                    "HORA DE SALIDA": st.column_config.TextColumn("HORA SALIDA", width=110),
                    "TIEMPO ADICIONAL": st.column_config.TextColumn("TIEMPO ADICIONAL", width=150),
                    "TOTAL HORAS": st.column_config.TextColumn("TOTAL HORAS", width=120),
                    "COMPENSADO": st.column_config.TextColumn("COMPENSADO", width=110),
                    "FECHA COMPENSACIÓN": st.column_config.TextColumn("FECHA COMP.", width=120),
                    "DESCRIPCIÓN": st.column_config.TextColumn("DESCRIPCIÓN", width=280),
                    "☑️ SELECCIONAR": st.column_config.CheckboxColumn("SELEC.", default=False, width=65),
                    "✏️ EDITAR": st.column_config.CheckboxColumn("EDIT.", default=False, width=60),
                    "🗑️ ELIMINAR": st.column_config.CheckboxColumn("ELIM.", default=False, width=60),
                },
            )

            # =================================================
            # EDICIÓN MASIVA DE VARIOS REGISTROS
            # =================================================
            seleccion_masiva = resultado_com[
                resultado_com["☑️ SELECCIONAR"] == True
            ].copy()

            if not seleccion_masiva.empty:
                ids_masivos = [
                    int(x) for x in seleccion_masiva["_ID_COMISION_"].tolist()
                ]

                with st.expander(
                    f"🛠️ EDITAR {len(ids_masivos)} REGISTRO(S) SELECCIONADO(S)",
                    expanded=True,
                ):
                    st.info(
                        "Selecciona varios registros en la columna **SELEC.** "
                        "y aplica el mismo cambio a todos. Las horas de ingreso y salida "
                        "no serán modificadas."
                    )

                    bm1, bm2 = st.columns(2)

                    with bm1:
                        opciones_adicional_masivo = [
                            "No cambiar",
                            "Sin aumento",
                            "30 minutos",
                            "1 hora",
                            "1 hora 30 min",
                            "2 horas",
                        ]
                        adicional_masivo = st.selectbox(
                            "➕ TIEMPO ADICIONAL",
                            opciones_adicional_masivo,
                            key=f"adicional_masivo_{st.session_state.comision_tabla_version}",
                        )

                    with bm2:
                        estado_comp_masivo = st.selectbox(
                            "🧾 COMPENSACIÓN",
                            [
                                "No cambiar",
                                "Marcar como compensadas",
                                "Dejar pendientes",
                            ],
                            key=f"estado_comp_masivo_{st.session_state.comision_tabla_version}",
                        )

                    fecha_comp_masiva = None
                    if estado_comp_masivo == "Marcar como compensadas":
                        fecha_comp_masiva = st.date_input(
                            "📅 FECHA DE COMPENSACIÓN",
                            value=fecha_local_hoy(),
                            format="DD/MM/YYYY",
                            key=f"fecha_comp_masiva_{st.session_state.comision_tabla_version}",
                        )

                    bm_guardar, bm_cancelar = st.columns([3, 1])

                    with bm_guardar:
                        aplicar_masivo = st.button(
                            "💾 APLICAR CAMBIOS A LOS SELECCIONADOS",
                            type="primary",
                            use_container_width=True,
                            key=f"aplicar_masivo_{st.session_state.comision_tabla_version}",
                        )

                    with bm_cancelar:
                        cancelar_masivo = st.button(
                            "✕ CANCELAR",
                            use_container_width=True,
                            key=f"cancelar_masivo_{st.session_state.comision_tabla_version}",
                        )

                    if cancelar_masivo:
                        st.session_state.comision_tabla_version += 1
                        st.rerun()

                    if aplicar_masivo:
                        if (
                            adicional_masivo == "No cambiar"
                            and estado_comp_masivo == "No cambiar"
                        ):
                            st.warning("⚠️ Selecciona al menos un cambio para aplicar.")
                        else:
                            mapa_adicional = {
                                "Sin aumento": 0.0,
                                "30 minutos": 0.5,
                                "1 hora": 1.0,
                                "1 hora 30 min": 1.5,
                                "2 horas": 2.0,
                            }

                            actualizados = 0
                            sin_horario = 0
                            errores = []

                            for id_masivo in ids_masivos:
                                try:
                                    registro = df_com[
                                        df_com["_ID_COMISION_"] == id_masivo
                                    ]
                                    if registro.empty:
                                        continue

                                    reg_m = registro.iloc[0]
                                    hora_ing_m = reg_m.get("HORA DE INGRESO")
                                    hora_sal_m = reg_m.get("HORA DE SALIDA")
                                    tiene_horario_m = (
                                        str(hora_ing_m or "").strip() not in ("", "None", "nan")
                                        and str(hora_sal_m or "").strip() not in ("", "None", "nan")
                                    )

                                    fecha_m = convertir_fecha(reg_m.get("FECHA"))
                                    hora_ing_obj = hora_segura_m = None
                                    hora_sal_obj = None

                                    def _hora_segura_m(valor):
                                        if valor in (None, "", "nan", "NaT"):
                                            return None
                                        try:
                                            return datetime.strptime(str(valor)[:8], "%H:%M:%S").time()
                                        except Exception:
                                            try:
                                                return datetime.strptime(str(valor)[:5], "%H:%M").time()
                                            except Exception:
                                                return None

                                    hora_ing_obj = _hora_segura_m(hora_ing_m)
                                    hora_sal_obj = _hora_segura_m(hora_sal_m)

                                    base_m = calcular_total_horas(
                                        fecha_m or fecha_local_hoy(),
                                        hora_ing_obj,
                                        hora_sal_obj,
                                    )

                                    if adicional_masivo == "No cambiar":
                                        try:
                                            adicional_actual_m = float(
                                                reg_m.get("TIEMPO ADICIONAL", 0) or 0
                                            )
                                        except (TypeError, ValueError):
                                            adicional_actual_m = 0.0
                                    else:
                                        adicional_actual_m = mapa_adicional[adicional_masivo]

                                    # Sin ingreso y salida no se asigna tiempo adicional ni total.
                                    adicional_aplicado_m = (
                                        adicional_actual_m if base_m is not None else 0.0
                                    )
                                    total_m = (
                                        base_m + adicional_aplicado_m
                                        if base_m is not None
                                        else None
                                    )

                                    if base_m is None and adicional_masivo != "No cambiar":
                                        sin_horario += 1

                                    compensado_m = es_valor_compensado(
                                        reg_m.get("COMPENSADO", False)
                                    )
                                    fecha_comp_m = convertir_fecha(
                                        reg_m.get("FECHA COMPENSACIÓN")
                                    )

                                    if estado_comp_masivo == "Marcar como compensadas":
                                        if total_m is not None:
                                            compensado_m = True
                                            fecha_comp_m = fecha_comp_masiva
                                        else:
                                            # No se puede compensar un registro sin total.
                                            compensado_m = False
                                            fecha_comp_m = None
                                            sin_horario += 1
                                    elif estado_comp_masivo == "Dejar pendientes":
                                        compensado_m = False
                                        fecha_comp_m = None

                                    datos_m = {
                                        "OPERADOR": reg_m.get("OPERADOR", ""),
                                        "FECHA": (fecha_m or fecha_local_hoy()).isoformat(),
                                        "TIPO DE SALIDA": reg_m.get("TIPO DE SALIDA", ""),
                                        "HORA DE INGRESO": (
                                            hora_ing_obj.strftime("%H:%M:%S")
                                            if hora_ing_obj is not None else None
                                        ),
                                        "HORA DE SALIDA": (
                                            hora_sal_obj.strftime("%H:%M:%S")
                                            if hora_sal_obj is not None else None
                                        ),
                                        "TIEMPO ADICIONAL": adicional_aplicado_m,
                                        "TOTAL HORAS": total_m,
                                        "COMPENSADO": compensado_m,
                                        "FECHA COMPENSACIÓN": (
                                            fecha_comp_m.isoformat()
                                            if fecha_comp_m else None
                                        ),
                                        "DESCRIPCIÓN": reg_m.get("DESCRIPCIÓN", ""),
                                    }

                                    actualizar_comision_supabase(datos_m, id_masivo)
                                    actualizados += 1

                                except Exception as e:
                                    errores.append(f"ID {id_masivo}: {e}")

                            st.session_state.comisiones = cargar_comisiones_supabase()
                            st.session_state.comision_tabla_version += 1

                            if actualizados:
                                st.success(
                                    f"✅ Se actualizaron {actualizados} registro(s) correctamente."
                                )
                            if sin_horario:
                                st.warning(
                                    f"⚠️ {sin_horario} registro(s) no tienen hora de ingreso y salida completas; "
                                    "se dejaron sin tiempo adicional/total."
                                )
                            if errores:
                                st.error(
                                    "❌ Algunos registros no pudieron actualizarse: "
                                    + " | ".join(errores)
                                )

                            st.rerun()

            # Eliminación segura desde la misma tabla.
            seleccion_eliminar = resultado_com[resultado_com["🗑️ ELIMINAR"] == True]
            if not seleccion_eliminar.empty:
                id_eliminar = int(seleccion_eliminar.iloc[0]["_ID_COMISION_"])
                registro_eliminar = df_com[df_com["_ID_COMISION_"] == id_eliminar]
                if not registro_eliminar.empty:
                    reg_elim = registro_eliminar.iloc[0]
                    st.warning(
                        f"⚠️ Vas a eliminar la salida de **{reg_elim.get('OPERADOR', '')}** "
                        f"del **{convertir_fecha(reg_elim.get('FECHA')).strftime('%d/%m/%Y') if convertir_fecha(reg_elim.get('FECHA')) else reg_elim.get('FECHA', '')}**. "
                        "Esta acción no se puede deshacer."
                    )
                    c_el1, c_el2 = st.columns([1, 1])
                    with c_el1:
                        confirmar_eliminar = st.button(
                            "🗑️ CONFIRMAR ELIMINACIÓN",
                            type="primary",
                            use_container_width=True,
                            key=f"confirmar_eliminar_comision_{id_eliminar}",
                        )
                    with c_el2:
                        cancelar_eliminar = st.button(
                            "✕ CANCELAR",
                            use_container_width=True,
                            key=f"cancelar_eliminar_comision_{id_eliminar}",
                        )
                    if cancelar_eliminar:
                        st.session_state.comision_tabla_version += 1
                        st.rerun()
                    if confirmar_eliminar:
                        try:
                            eliminar_comision_supabase(id_eliminar)
                            st.session_state.comisiones = cargar_comisiones_supabase()
                            st.session_state.comision_id_editar_inline = None
                            st.session_state.comision_tabla_version += 1
                            st.success("✅ Registro eliminado correctamente.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ No se pudo eliminar el registro: {e}")

            seleccion = resultado_com[resultado_com["✏️ EDITAR"] == True]
            if not seleccion.empty:
                id_seleccionado = int(seleccion.iloc[0]["_ID_COMISION_"])
                if st.session_state.get("comision_id_editar_inline") != id_seleccionado:
                    st.session_state.comision_id_editar_inline = id_seleccionado
                    st.session_state.comision_tabla_version += 1
                    st.rerun()

            # =================================================
            # FORMULARIO DE EDICIÓN DEL REGISTRO SELECCIONADO
            # =================================================
            id_editar = st.session_state.get("comision_id_editar_inline")
            if id_editar is not None:
                registro_editar = df_com[df_com["_ID_COMISION_"] == id_editar]
                if not registro_editar.empty:
                    reg = registro_editar.iloc[0]

                    def hora_segura(valor):
                        if valor in (None, "", "nan", "NaT"):
                            return None
                        try:
                            return datetime.strptime(str(valor)[:8], "%H:%M:%S").time()
                        except Exception:
                            try:
                                return datetime.strptime(str(valor)[:5], "%H:%M").time()
                            except Exception:
                                return None

                    fecha_edit = convertir_fecha(reg.get("FECHA")) or fecha_local_hoy()
                    hora_ing_edit = hora_segura(reg.get("HORA DE INGRESO"))
                    hora_sal_edit = hora_segura(reg.get("HORA DE SALIDA"))
                    try:
                        tiempo_adicional_edit = float(reg.get("TIEMPO ADICIONAL", 0) or 0)
                    except (TypeError, ValueError):
                        tiempo_adicional_edit = 0.0
                    compensado_edit = es_valor_compensado(reg.get("COMPENSADO", False))
                    fecha_comp_edit = convertir_fecha(reg.get("FECHA COMPENSACIÓN"))

                    with st.expander("✏️ EDITAR REGISTRO SELECCIONADO", expanded=True):
                        with st.form(
                            f"form_editar_comision_{id_editar}_{st.session_state.comision_tabla_version}",
                            clear_on_submit=False,
                        ):
                            e1, e2, e3 = st.columns(3)

                            with e1:
                                op_actual = str(reg.get("OPERADOR", ""))
                                operadores_ed = CAT_OPERADORES_COMISION
                                operador_edit = st.selectbox(
                                    "👤 OPERADOR",
                                    operadores_ed,
                                    index=operadores_ed.index(op_actual) if op_actual in operadores_ed else 0,
                                )
                                fecha_edit = st.date_input(
                                    "📅 FECHA",
                                    value=fecha_edit,
                                    format="DD/MM/YYYY",
                                )

                            with e2:
                                tipo_actual = str(reg.get("TIPO DE SALIDA", ""))
                                tipo_edit = st.selectbox(
                                    "🚶 TIPO DE SALIDA",
                                    CAT_TIPOS_SALIDA_COMISION,
                                    index=CAT_TIPOS_SALIDA_COMISION.index(tipo_actual)
                                    if tipo_actual in CAT_TIPOS_SALIDA_COMISION else 0,
                                )
                                tiene_ing = st.checkbox(
                                    "Registrar HORA DE INGRESO",
                                    value=hora_ing_edit is not None,
                                )
                                hora_ingreso_edit = (
                                    st.time_input(
                                        "🕘 HORA DE INGRESO",
                                        value=hora_ing_edit or hora_actual_local(),
                                        step=60,
                                    )
                                    if tiene_ing else None
                                )

                            with e3:
                                tiene_sal = st.checkbox(
                                    "Registrar HORA DE SALIDA",
                                    value=hora_sal_edit is not None,
                                )
                                hora_salida_edit = (
                                    st.time_input(
                                        "🕘 HORA DE SALIDA",
                                        value=hora_sal_edit or hora_actual_local(),
                                        step=60,
                                    )
                                    if tiene_sal else None
                                )

                            opciones_adicional_edit = list(TIEMPOS_ADICIONALES.keys())
                            etiqueta_adicional_edit = etiqueta_tiempo_adicional(tiempo_adicional_edit)
                            if etiqueta_adicional_edit not in opciones_adicional_edit:
                                etiqueta_adicional_edit = "Sin aumento"
                            tiempo_adicional_label_edit = st.selectbox(
                                "➕ TIEMPO ADICIONAL PARA COMPENSACIÓN",
                                opciones_adicional_edit,
                                index=opciones_adicional_edit.index(etiqueta_adicional_edit),
                                help="Suma tiempo de compensación sin modificar las horas registradas.",
                            )
                            tiempo_adicional_edit = TIEMPOS_ADICIONALES[tiempo_adicional_label_edit]

                            descripcion_edit = st.text_area(
                                "📝 DESCRIPCIÓN",
                                value=str(reg.get("DESCRIPCIÓN", "")),
                            )

                            horas_base_edit = calcular_total_horas(
                                fecha_edit,
                                hora_ingreso_edit,
                                hora_salida_edit,
                            )
                            tiempo_adicional_edit_aplicado = (
                                tiempo_adicional_edit if horas_base_edit is not None else 0.0
                            )
                            total_edit = (
                                horas_base_edit + tiempo_adicional_edit_aplicado
                                if horas_base_edit is not None else None
                            )
                            if total_edit is not None:
                                st.info(
                                    f"⏱️ TOTAL HORAS: {formatear_horas_minutos(total_edit)} "
                                    f"(base: {formatear_horas_minutos(horas_base_edit)} + adicional: {formatear_horas_minutos(tiempo_adicional_edit_aplicado)})"
                                )
                            else:
                                st.caption("⏱️ TOTAL HORAS: vacío hasta que se registren ambas horas.")

                            st.markdown("#### 🧾 ESTADO DE COMPENSACIÓN")
                            compensado_edit = st.checkbox(
                                "✅ Marcar estas horas como COMPENSADAS",
                                value=compensado_edit,
                            )

                            if compensado_edit:
                                fecha_comp_edit = st.date_input(
                                    "📅 FECHA DE COMPENSACIÓN",
                                    value=fecha_comp_edit or fecha_local_hoy(),
                                    format="DD/MM/YYYY",
                                )
                            else:
                                fecha_comp_edit = None
                                st.caption("⏳ Este registro quedará pendiente de compensación.")

                            g1, g2 = st.columns([4, 1])
                            with g1:
                                actualizar = st.form_submit_button(
                                    "💾 ACTUALIZAR REGISTRO",
                                    use_container_width=True,
                                )
                            with g2:
                                cancelar_ed = st.form_submit_button(
                                    "✕ CANCELAR",
                                    type="primary",
                                    use_container_width=True,
                                )

                        if cancelar_ed:
                            st.session_state.comision_id_editar_inline = None
                            st.session_state.comision_tabla_version += 1
                            st.rerun()

                        if actualizar:
                            if operador_edit == CAT_OPERADORES_COMISION[0]:
                                st.error("❌ Selecciona el OPERADOR.")
                            elif tipo_edit == CAT_TIPOS_SALIDA_COMISION[0]:
                                st.error("❌ Selecciona el TIPO DE SALIDA.")
                            elif not descripcion_edit.strip():
                                st.error("❌ Ingresa una DESCRIPCIÓN de la comisión.")
                            elif compensado_edit and total_edit is None:
                                st.error("❌ Para marcar como compensadas, primero debes registrar HORA DE INGRESO y HORA DE SALIDA.")
                            else:
                                try:
                                    datos_edit = {
                                        "OPERADOR": operador_edit,
                                        "FECHA": fecha_edit.isoformat(),
                                        "TIPO DE SALIDA": tipo_edit,
                                        "HORA DE INGRESO": hora_ingreso_edit.strftime("%H:%M:%S") if hora_ingreso_edit is not None else None,
                                        "HORA DE SALIDA": hora_salida_edit.strftime("%H:%M:%S") if hora_salida_edit is not None else None,
                                        "TIEMPO ADICIONAL": tiempo_adicional_edit_aplicado,
                                        "TOTAL HORAS": total_edit,
                                        "COMPENSADO": compensado_edit,
                                        "FECHA COMPENSACIÓN": fecha_comp_edit.isoformat() if fecha_comp_edit else None,
                                        "DESCRIPCIÓN": descripcion_edit.strip(),
                                    }
                                    actualizar_comision_supabase(datos_edit, id_editar)
                                    st.session_state.comisiones = cargar_comisiones_supabase()
                                    st.session_state.comision_id_editar_inline = None
                                    st.session_state.comision_tabla_version += 1
                                    st.success("✅ Registro actualizado correctamente.")
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"❌ No se pudo actualizar el registro: {e}")

            # IMPORTANTE: descargar df_filtrado, NO df_com.
            st.download_button(
                "📥 DESCARGAR FILTRADO EN EXCEL",
                data=excel_bytes_comisiones(df_filtrado),
                file_name=f"SALIDAS_COMISION_FILTRADO_{fecha_local_hoy().isoformat()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

    except Exception as e:
        st.error(f"❌ No se pudo cargar el historial de salidas: {e}")

    st.stop()



# =========================================================
# PAGINA CONTROL DE CAJAS Y CINTILLOS
# =========================================================
if st.session_state.page == "cajas":
    st.subheader("📦 CONTROL DE CAJAS Y CINTILLOS")
    st.caption("Control de ingresos, stock actual y salidas de cajas y cintillos.")
    if st.button("⬅ VOLVER A SOLICITUDES", key="volver_cajas"):
        st.session_state.page = "solicitudes"
        st.rerun()

    try:
        stock = obtener_stock_actual()
        m1, m2, m3 = st.columns(3)
        m1.metric("📦 CAJAS NUEVAS", stock["CAJAS NUEVAS"])
        m2.metric("♻️ CAJAS REUTILIZADAS", stock["CAJAS REUTILIZADAS"])
        m3.metric("🏷️ CINTILLOS", stock["CINTILLOS"])
    except Exception as e:
        st.error(f"❌ No se pudo calcular el stock: {e}")
        st.stop()

    if es_admin():
        # El título cambia internamente después de CANCELAR/GUARDAR para que
        # Streamlit reconstruya el expander y vuelva a quedar cerrado.
        st.session_state.setdefault("stock_form_version", 0)
        st.session_state.setdefault("stock_expander_version", 0)

        titulo_stock = (
            "➕ INGRESAR LLEGADA DE STOCK"
            if st.session_state.stock_expander_version % 2 == 0
            else "➕ INGRESAR LLEGADA DE STOCK\u2060"
        )

        with st.expander(titulo_stock, expanded=False):
            st.caption("🔴 Solo ADMIN puede aumentar o ajustar el stock.")

            with st.form(
                f"stock_form_{st.session_state.stock_form_version}",
                clear_on_submit=False
            ):
                a, b = st.columns(2)

                with a:
                    fecha_ingreso_stock = st.date_input(
                        "📅 Fecha de llegada",
                        value=fecha_local_hoy(),
                        format="DD/MM/YYYY"
                    )
                    cajas_nuevas_in = st.number_input(
                        "📦 Cajas nuevas recibidas",
                        min_value=0,
                        step=1,
                        value=0
                    )
                    cajas_reutilizadas_in = st.number_input(
                        "♻️ Cajas reutilizadas ingresadas",
                        min_value=0,
                        step=1,
                        value=0
                    )

                with b:
                    cintillos_in = st.number_input(
                        "🏷️ Cintillos recibidos",
                        min_value=0,
                        step=1,
                        value=0
                    )
                    proveedor = st.text_input("🚚 Proveedor / procedencia")
                    obs_ingreso = st.text_area("📝 Observaciones")

                col_guardar, col_cancelar = st.columns([4, 1])

                with col_guardar:
                    guardar_stock = st.form_submit_button(
                        "💾 GUARDAR INGRESO DE STOCK",
                        use_container_width=True
                    )

                with col_cancelar:
                    cancelar_stock = st.form_submit_button(
                        "✕ CANCELAR",
                        type="primary",
                        use_container_width=True
                    )

            # CANCELAR: no guarda nada, limpia los widgets y repliega el expander.
            if cancelar_stock:
                st.session_state.stock_form_version += 1
                st.session_state.stock_expander_version += 1
                st.rerun()

            if guardar_stock:
                total_ingreso = (
                    int(cajas_nuevas_in)
                    + int(cajas_reutilizadas_in)
                    + int(cintillos_in)
                )

                if total_ingreso <= 0:
                    st.error("Ingresa al menos una cantidad mayor a cero.")
                else:
                    guardar_movimiento_stock(
                        fecha_ingreso_stock,
                        "INGRESO",
                        cajas_nuevas_in,
                        cajas_reutilizadas_in,
                        cintillos_in,
                        proveedor,
                        obs_ingreso
                    )
                    st.session_state.cajas = cargar_cajas_supabase()
                    st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
                    st.session_state.stock_form_version += 1
                    st.session_state.stock_expander_version += 1
                    st.success("✅ Ingreso de stock registrado correctamente.")
                    st.rerun()
    else:
        st.info("🔒 Solo el usuario ADMIN puede ingresar o modificar el stock.")

    if puede_modificar():
        # Registro de salida como sección desplegable, igual que "Ingresar llegada de stock".
        # Se usa una versión invisible del título para forzar que Streamlit reconstruya
        # el expander cuando se presiona CANCELAR o se guarda el registro.
        st.session_state.setdefault("mostrar_registro_salida", False)
        st.session_state.setdefault("salida_expander_version", 0)

        titulo_salida = (
            "📝 REGISTRAR SALIDA DE CAJAS Y CINTILLOS"
            if st.session_state.salida_expander_version % 2 == 0
            else "📝 REGISTRAR SALIDA DE CAJAS Y CINTILLOS\u2060"
        )

        with st.expander(
            titulo_salida,
            expanded=st.session_state.mostrar_registro_salida,
        ):
            with st.form(f"salida_cajas_form_{st.session_state.caja_form_version}"):
                a,b = st.columns(2)
                with a:
                    solicitante = st.text_input("SOLICITANTE")
                    cliente_caja = st.selectbox("CLIENTE", CAT_CLIENTES[1:])
                    agencia_caja = st.selectbox(
                        "AGENCIA",
                        CAT_AGENCIAS[1:],
                        index=None,
                        placeholder="Seleccione una agencia"
                    )
                    fecha_salida = st.date_input("FECHA DE SALIDA", value=fecha_local_hoy(), format="DD/MM/YYYY")
                    cajas_solicitadas = st.number_input("CANTIDAD DE CAJAS SOLICITADAS", min_value=0, step=1, value=0)
                with b:
                    nro_wo_caja = st.text_input("NRO. WORKORDER")
                    cajas_reutilizadas = st.number_input("CAJAS REUTILIZADAS", min_value=0, step=1, value=0)
                    cintillos_salida = st.number_input("CANTIDAD DE CINTILLOS", min_value=0, step=1, value=0)
                    observaciones_caja = st.text_area("OBSERVACIONES")

                x, y = st.columns([4, 1])
                with x:
                    guardar_salida = st.form_submit_button("📤 REGISTRAR SALIDA", use_container_width=True)
                with y:
                    cancelar_salida = st.form_submit_button("✕ CANCELAR", type="primary", use_container_width=True)

            if cancelar_salida:
                # Limpiar el formulario y cerrar visualmente el expander.
                st.session_state.caja_form_version += 1
                st.session_state.mostrar_registro_salida = False
                st.session_state.salida_expander_version += 1
                st.rerun()

            if guardar_salida:
                cajas_nuevas_salida = int(cajas_solicitadas) - int(cajas_reutilizadas)
                if not solicitante.strip() or not agencia_caja or (int(cajas_solicitadas) <= 0 and int(cintillos_salida) <= 0):
                    st.error("Completa SOLICITANTE y AGENCIA, e ingresa al menos una cantidad mayor a cero de CAJAS o CINTILLOS.")
                elif cajas_nuevas_salida < 0:
                    st.error("Las cajas reutilizadas no pueden ser mayores que las cajas solicitadas.")
                elif cajas_nuevas_salida > stock["CAJAS NUEVAS"] or int(cajas_reutilizadas) > stock["CAJAS REUTILIZADAS"] or int(cintillos_salida) > stock["CINTILLOS"]:
                    st.error("❌ Stock insuficiente para registrar esta salida.")
                else:
                    datos_salida = {"SOLICITANTE": solicitante.strip(), "CLIENTE": cliente_caja, "AGENCIA": agencia_caja.strip(), "FECHA DE SALIDA": fecha_salida.isoformat(), "CANTIDAD DE CAJAS SOLICITADAS": int(cajas_solicitadas), "NRO. WORKORDER": nro_wo_caja.strip(), "CAJAS REUTILIZADAS": int(cajas_reutilizadas), "CANTIDAD DE CINTILLOS": int(cintillos_salida), "OBSERVACIONES": observaciones_caja.strip()}
                    guardar_salida_cajas_supabase(datos_salida)
                    guardar_movimiento_stock(fecha_salida, "SALIDA", -cajas_nuevas_salida, -int(cajas_reutilizadas), -int(cintillos_salida), "", f"Salida WO: {nro_wo_caja.strip()}")
                    st.session_state.cajas = cargar_cajas_supabase()
                    st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
                    st.session_state.caja_form_version += 1
                    st.session_state.mostrar_registro_salida = False
                    st.session_state.salida_expander_version += 1
                    st.success("✅ Salida registrada y stock actualizado automáticamente.")
                    st.rerun()

    else:
        st.info("👁️ MODO INVITADO — Puedes consultar el stock y el historial de salidas, pero no registrar salidas.")
    st.markdown("### 📋 Historial de salidas")
    cajas_df = st.session_state.cajas if st.session_state.cajas is not None else cargar_cajas_supabase()

    # Solo el usuario ADMIN puede filtrar y descargar el historial completo.
    if es_admin():
        st.session_state.setdefault("mostrar_filtro_historial_salidas", False)

        b1, b2 = st.columns([1.2, 1.2])
        with b1:
            if st.button(
                "🔎 FILTRAR HISTORIAL",
                use_container_width=True,
                key="btn_filtro_historial_salidas",
            ):
                # No forzamos un segundo rerun: así no se vuelve a abrir
                # accidentalmente el formulario de REGISTRAR SALIDA.
                st.session_state.mostrar_filtro_historial_salidas = (
                    not st.session_state.mostrar_filtro_historial_salidas
                )

        # El historial descargado será exactamente el que se muestra después
        # de aplicar los filtros seleccionados.
        tabla_filtrada = cajas_df.copy()

        # Estas variables deben existir aunque el panel de filtros esté cerrado.
        # Antes, el informe mensual las utilizaba desde fuera del bloque de
        # filtros y Streamlit producía NameError cuando el filtro estaba cerrado.
        meses_es = {
            1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL",
            5: "MAYO", 6: "JUNIO", 7: "JULIO", 8: "AGOSTO",
            9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE",
        }
        periodo_seleccionado = None
        mes_historial = "TODOS"
        cliente_historial = "TODOS"

        if st.session_state.mostrar_filtro_historial_salidas:
            st.markdown("#### 🔎 Filtros del historial de salidas")

            # Opciones del filtro por MES en español, por ejemplo:
            # SEPTIEMBRE 2026.
            if cajas_df.empty:
                st.info("No hay registros de salidas para filtrar.")
            else:
                fechas_hist = pd.to_datetime(
                    cajas_df["FECHA DE SALIDA"], errors="coerce"
                ).dropna()

                if fechas_hist.empty:
                    opciones_mes = [("TODOS", None)]
                else:
                    periodos = sorted(
                        fechas_hist.dt.to_period("M").unique().tolist(),
                        reverse=True,
                    )
                    opciones_mes = [("TODOS", None)] + [
                        (f"{meses_es[int(periodo.month)]} {int(periodo.year)}", periodo)
                        for periodo in periodos
                    ]

                # ======================================================
                # FILTROS INDEPENDIENTES Y COMBINADOS
                # Cada selector se alimenta del historial COMPLETO.
                # Así:
                #   - MES funciona por separado.
                #   - CLIENTE funciona por separado.
                #   - Si se seleccionan ambos, se cumplen AMBAS condiciones.
                # ======================================================
                f1, f2 = st.columns(2)

                with f1:
                    etiquetas_mes = [etiqueta for etiqueta, _ in opciones_mes]
                    mes_historial = st.selectbox(
                        "📅 FILTRAR POR MES",
                        etiquetas_mes,
                        key="filtro_historial_mes",
                    )

                with f2:
                    clientes_historial = sorted(
                        [
                            str(cliente).strip()
                            for cliente in cajas_df["CLIENTE"].dropna().unique().tolist()
                            if str(cliente).strip()
                        ]
                    )
                    opciones_cliente = ["TODOS"] + clientes_historial
                    cliente_historial = st.selectbox(
                        "👤 FILTRAR POR CLIENTE",
                        opciones_cliente,
                        key="filtro_historial_cliente",
                    )

                # Empezamos siempre desde el historial completo.
                # De esta forma los filtros son independientes.
                tabla_filtrada = cajas_df.copy()

                # 1) FILTRO POR MES (si se seleccionó un mes).
                periodo_seleccionado = dict(opciones_mes).get(mes_historial)
                if periodo_seleccionado is not None:
                    periodos_tabla = pd.to_datetime(
                        tabla_filtrada["FECHA DE SALIDA"], errors="coerce"
                    ).dt.to_period("M")
                    tabla_filtrada = tabla_filtrada[
                        periodos_tabla == periodo_seleccionado
                    ]

                # 2) FILTRO POR CLIENTE (si se seleccionó un cliente).
                # Si también hay un mes seleccionado, este filtro se aplica
                # sobre el resultado anterior, cumpliendo AMBAS condiciones.
                if cliente_historial != "TODOS":
                    tabla_filtrada = tabla_filtrada[
                        tabla_filtrada["CLIENTE"].astype(str).str.strip() == cliente_historial
                    ]

                # Indicamos claramente qué filtro está activo.
                filtros_activos = []
                if mes_historial != "TODOS":
                    filtros_activos.append(f"Mes: {mes_historial}")
                if cliente_historial != "TODOS":
                    filtros_activos.append(f"Cliente: {cliente_historial}")

                if filtros_activos:
                    if len(filtros_activos) == 2:
                        st.caption(
                            "Mostrando registros que cumplen ambos filtros: "
                            + "  •  ".join(filtros_activos)
                        )
                    else:
                        st.caption(
                            "Mostrando registros filtrados por: "
                            + "  •  ".join(filtros_activos)
                        )
                else:
                    st.caption("Mostrando todo el historial de salidas.")

                # Resumen de cantidades correspondientes exactamente al filtro aplicado.
                total_cajas_filtradas = int(
                    pd.to_numeric(
                        tabla_filtrada["CANTIDAD DE CAJAS SOLICITADAS"],
                        errors="coerce",
                    ).fillna(0).sum()
                ) if "CANTIDAD DE CAJAS SOLICITADAS" in tabla_filtrada.columns else 0

                total_cintillos_filtrados = int(
                    pd.to_numeric(
                        tabla_filtrada["CANTIDAD DE CINTILLOS"],
                        errors="coerce",
                    ).fillna(0).sum()
                ) if "CANTIDAD DE CINTILLOS" in tabla_filtrada.columns else 0

                st.markdown("#### 📊 Totales del filtro")
                t1, t2 = st.columns(2)
                with t1:
                    st.metric("📦 TOTAL DE CAJAS SOLICITADAS", f"{total_cajas_filtradas:,}")
                with t2:
                    st.metric("🏷️ TOTAL DE CINTILLOS SOLICITADOS", f"{total_cintillos_filtrados:,}")

        with b2:
            nombre_archivo = (
                "historial_salidas_"
                + datetime.now(ZONA_HORARIA_APP).strftime("%Y%m%d_%H%M%S")
                + ".xlsx"
            )
            st.download_button(
                "⬇️ DESCARGAR HISTORIAL",
                data=excel_bytes_historial_salidas(tabla_filtrada),
                file_name=nombre_archivo,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="descargar_historial_salidas",
                disabled=tabla_filtrada.empty,
            )
            # ==========================================================
        # INFORME MENSUAL PARA PRESENTAR A JEFATURA
        # El informe se genera por MES. Para evitar reportes
        # incompletos, se debe seleccionar un mes distinto de TODOS.
        # El filtro de cliente puede combinarse con el mes.
        # ==========================================================
        informe_mensual_disponible = (
            periodo_seleccionado is not None and not tabla_filtrada.empty
        )

        if periodo_seleccionado is not None:
            fecha_desde_informe = periodo_seleccionado.start_time
            fecha_hasta_informe = periodo_seleccionado.end_time
            nombre_mes_informe = (
                f"{meses_es[int(periodo_seleccionado.month)]}_"
                f"{int(periodo_seleccionado.year)}"
            )
            titulo_informe = (
                f"Mes: {meses_es[int(periodo_seleccionado.month)]} "
                f"{int(periodo_seleccionado.year)}"
            )
            if cliente_historial != "TODOS":
                titulo_informe += f" | Cliente: {cliente_historial}"
        else:
            fecha_desde_informe = None
            fecha_hasta_informe = None
            nombre_mes_informe = "SELECCIONAR_MES"
            titulo_informe = "Seleccione un mes para generar el informe mensual"

        # ==========================================================
        # DATOS DEL KARDEX POR CLIENTE
        # Un solo código/precio por cliente se aplica a TODOS sus registros.
        # BASE y TAPA tienen precios independientes.
        # ==========================================================
        datos_clientes = preparar_editor_kardex_clientes(tabla_filtrada)
        firma_clientes = "|".join(datos_clientes["CLIENTE"].astype(str).tolist()) if not datos_clientes.empty else ""
        firma_clientes = hashlib.md5(f"{firma_clientes}|{nombre_mes_informe}|{cliente_historial}".encode("utf-8")).hexdigest()

        datos_clientes_editor = datos_clientes.copy()
        if st.session_state.get("kardex_editor_firma") != firma_clientes:
            st.session_state.kardex_editor_firma = firma_clientes
            st.session_state.kardex_clientes_editados = datos_clientes.copy()
        elif st.session_state.kardex_clientes_editados is not None:
            datos_clientes_editor = st.session_state.kardex_clientes_editados.copy()

        if informe_mensual_disponible:
            st.markdown("#### 💰 DATOS PARA GENERAR EL KARDEX")
            st.caption(
                "PRECIO BASE, PRECIO TAPA y TIPO DE CAMBIO son iguales para todos los clientes. "
                "En la tabla solamente cambia el COD. CENTRO COSTO de cada cliente."
            )

            # ======================================================
            # VALORES GENERALES DEL KARDEX
            # Se ingresan una sola vez y se aplican a TODOS los clientes.
            # ======================================================
            cfg_actual = st.session_state.kardex_clientes_editados
            if cfg_actual is None or cfg_actual.empty:
                cfg_actual = datos_clientes.copy()

            def _primer_valor_numerico(columna, defecto):
                try:
                    serie = pd.to_numeric(cfg_actual.get(columna, pd.Series(dtype=float)), errors="coerce").dropna()
                    if not serie.empty and float(serie.iloc[0]) > 0:
                        return float(serie.iloc[0])
                except Exception:
                    pass
                return float(defecto)

            k1, k2, k3 = st.columns(3)
            with k1:
                precio_base_global = st.number_input(
                    "PRECIO BASE",
                    min_value=0.0,
                    step=0.01,
                    value=_primer_valor_numerico("PRECIO BASE", 0.0),
                    format="%.2f",
                    key=f"kardex_precio_base_{firma_clientes}",
                    help="Se aplica automáticamente a todos los clientes del Kardex.",
                )
            with k2:
                precio_tapa_global = st.number_input(
                    "PRECIO TAPA",
                    min_value=0.0,
                    step=0.01,
                    value=_primer_valor_numerico("PRECIO TAPA", 0.0),
                    format="%.2f",
                    key=f"kardex_precio_tapa_{firma_clientes}",
                    help="Se aplica automáticamente a todos los clientes del Kardex.",
                )
            with k3:
                tipo_cambio_global = st.number_input(
                    "TIPO DE CAMBIO",
                    min_value=0.0,
                    step=0.01,
                    value=_primer_valor_numerico("TIPO DE CAMBIO", 1.0),
                    format="%.2f",
                    key=f"kardex_tipo_cambio_{firma_clientes}",
                    help="Se aplica automáticamente a todos los clientes del Kardex.",
                )

            # ======================================================
            # TABLA POR CLIENTE
            # Solo el COD. CENTRO COSTO es diferente por cliente.
            # La tabla usa todo el ancho disponible.
            # ======================================================
            tabla_codigos = datos_clientes_editor[["CLIENTE", "COD. CENTRO COSTO"]].copy()
            tabla_codigos["COD. CENTRO COSTO"] = tabla_codigos["COD. CENTRO COSTO"].fillna("").astype(str).str.strip()

            editor_codigos = st.data_editor(
                tabla_codigos,
                hide_index=True,
                use_container_width=True,
                num_rows="fixed",
                column_config={
                    "CLIENTE": st.column_config.TextColumn("CLIENTE", disabled=True, width="large"),
                    "COD. CENTRO COSTO": st.column_config.TextColumn(
                        "COD. CENTRO COSTO",
                        help="Código correspondiente a cada cliente. Es el único dato que varía por cliente.",
                        width="large",
                    ),
                },
                key=f"editor_kardex_codigos_{firma_clientes}",
            )

            editor_codigos["CLIENTE"] = editor_codigos["CLIENTE"].fillna("").astype(str).str.strip()
            editor_codigos["COD. CENTRO COSTO"] = editor_codigos["COD. CENTRO COSTO"].fillna("").astype(str).str.strip()
            editor_clientes = editor_codigos.copy()
            editor_clientes["PRECIO BASE"] = float(precio_base_global)
            editor_clientes["PRECIO TAPA"] = float(precio_tapa_global)
            editor_clientes["TIPO DE CAMBIO"] = float(tipo_cambio_global)
            editor_clientes = editor_clientes[["CLIENTE", "COD. CENTRO COSTO", "PRECIO BASE", "PRECIO TAPA", "TIPO DE CAMBIO"]]
            st.session_state.kardex_clientes_editados = editor_clientes.copy()

            # Mostrar cuántos registros serán afectados por cada cliente.
            conteo_clientes = tabla_filtrada["CLIENTE"].fillna("").astype(str).str.strip().value_counts()
            vista_clientes = editor_clientes.copy()
            vista_clientes["REGISTROS"] = vista_clientes["CLIENTE"].map(conteo_clientes).fillna(0).astype(int)
            vista_clientes = vista_clientes[["CLIENTE", "REGISTROS", "COD. CENTRO COSTO", "PRECIO BASE", "PRECIO TAPA", "TIPO DE CAMBIO"]]

            st.caption("📌 El código y precios de cada cliente se aplicarán automáticamente a todos los registros mostrados en este mes.")

            # Convertimos la configuración por cliente en datos por cada salida.
            st.session_state.kardex_datos_editados = construir_kardex_por_registro(
                tabla_filtrada, editor_clientes
            )
            datos_aplicados = st.session_state.kardex_datos_editados
            total_base_preview = float(pd.to_numeric(datos_aplicados.get("TOTAL BASE", 0), errors="coerce").fillna(0).sum()) if not datos_aplicados.empty else 0.0
            total_tapa_preview = float(pd.to_numeric(datos_aplicados.get("TOTAL TAPA", 0), errors="coerce").fillna(0).sum()) if not datos_aplicados.empty else 0.0

            p1, p2, p3 = st.columns(3)
            with p1:
                st.metric("📦 VALOR BASE", f"Bs {total_base_preview:,.2f}")
            with p2:
                st.metric("🧢 VALOR TAPA", f"Bs {total_tapa_preview:,.2f}")
            with p3:
                st.metric("💰 VALOR TOTAL", f"Bs {total_base_preview + total_tapa_preview:,.2f}")

        nombre_informe = (
            "informe_kardex_cajas_cintillos_"
            + nombre_mes_informe
            + "_"
            + datetime.now(ZONA_HORARIA_APP).strftime("%Y%m%d_%H%M%S")
            + ".xlsx"
        )

        # Validacion: no permitir generar el Kardex si falta algun dato solicitado.
        datos_listos = False
        if informe_mensual_disponible and st.session_state.kardex_clientes_editados is not None:
            cfg = st.session_state.kardex_clientes_editados
            datos_listos = (
                not cfg.empty
                and cfg["COD. CENTRO COSTO"].astype(str).str.strip().ne("").all()
                and pd.to_numeric(cfg["PRECIO BASE"], errors="coerce").fillna(0).gt(0).all()
                and pd.to_numeric(cfg["PRECIO TAPA"], errors="coerce").fillna(0).gt(0).all()
                and pd.to_numeric(cfg["TIPO DE CAMBIO"], errors="coerce").fillna(0).gt(0).all()
            )

        st.download_button(
            "📊 DESCARGAR INFORME MENSUAL PARA PRESENTAR",
            data=excel_bytes_informe_cajas_cintillos(
                tabla_filtrada,
                st.session_state.movimientos_stock,
                titulo_informe,
                fecha_desde_informe,
                fecha_hasta_informe,
                st.session_state.kardex_datos_editados if informe_mensual_disponible else None,
            ),
            file_name=nombre_informe,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="descargar_informe_kardex_cajas",
            disabled=not (informe_mensual_disponible and datos_listos),
            help=(
                "Seleccione un mes y complete los datos del Kardex antes de descargar el informe."
            ),
        )
        if informe_mensual_disponible and not datos_listos:
            st.warning("⚠️ Complete COD. CENTRO COSTO, PRECIO BASE, PRECIO TAPA y TIPO DE CAMBIO para todos los clientes antes de generar el informe.")

        # La edición se selecciona directamente desde la columna ✏️ EDITAR de la tabla.
        st.session_state.setdefault("salida_id_editar_inline", None)
        salida_id_editar = st.session_state.get("salida_id_editar_inline")

        if salida_id_editar is not None:
            try:
                salida_id_editar = int(salida_id_editar)
                if salida_id_editar not in tabla_filtrada["_ID_CAJA_"].astype(int).tolist():
                    salida_id_editar = None
                    st.session_state.salida_id_editar_inline = None
            except Exception:
                salida_id_editar = None
                st.session_state.salida_id_editar_inline = None

        if salida_id_editar is not None:
            st.markdown("#### ✏️ Editar registro de salida")
            st.caption("🔴 Solo ADMIN puede editar. El stock se ajustará automáticamente según los cambios realizados.")

            registro_editar = tabla_filtrada[
                tabla_filtrada["_ID_CAJA_"].astype(int) == int(salida_id_editar)
            ].iloc[0].to_dict()

            cliente_actual = str(registro_editar.get("CLIENTE", "")).strip()
            clientes_edicion = list(CAT_CLIENTES[1:])
            if cliente_actual and cliente_actual not in clientes_edicion:
                clientes_edicion.append(cliente_actual)

            agencia_actual = str(registro_editar.get("AGENCIA", "")).strip()
            agencias_edicion = list(CAT_AGENCIAS[1:])
            if agencia_actual and agencia_actual not in agencias_edicion:
                agencias_edicion.append(agencia_actual)

            fecha_edicion = convertir_fecha(registro_editar.get("FECHA DE SALIDA")) or fecha_local_hoy()

            with st.form(f"editar_salida_form_{int(salida_id_editar)}"):
                e1, e2 = st.columns(2)
                with e1:
                    edit_solicitante = st.text_input(
                        "SOLICITANTE", value=str(registro_editar.get("SOLICITANTE", ""))
                    )
                    edit_cliente = st.selectbox(
                        "CLIENTE",
                        clientes_edicion,
                        index=clientes_edicion.index(cliente_actual) if cliente_actual in clientes_edicion else 0,
                    )
                    edit_agencia = st.selectbox(
                        "AGENCIA",
                        agencias_edicion,
                        index=agencias_edicion.index(agencia_actual) if agencia_actual in agencias_edicion else 0,
                    )
                    edit_fecha = st.date_input(
                        "FECHA DE SALIDA", value=fecha_edicion, format="DD/MM/YYYY"
                    )
                    edit_cajas = st.number_input(
                        "CANTIDAD DE CAJAS SOLICITADAS",
                        min_value=0,
                        step=1,
                        value=entero_seguro(registro_editar.get("CANTIDAD DE CAJAS SOLICITADAS", 0)),
                    )
                with e2:
                    edit_wo = st.text_input(
                        "NRO. WORKORDER", value=str(registro_editar.get("NRO. WORKORDER", ""))
                    )
                    edit_reutilizadas = st.number_input(
                        "CAJAS REUTILIZADAS",
                        min_value=0,
                        step=1,
                        value=entero_seguro(registro_editar.get("CAJAS REUTILIZADAS", 0)),
                    )
                    edit_cintillos = st.number_input(
                        "CANTIDAD DE CINTILLOS",
                        min_value=0,
                        step=1,
                        value=entero_seguro(registro_editar.get("CANTIDAD DE CINTILLOS", 0)),
                    )
                    edit_observaciones = st.text_area(
                        "OBSERVACIONES", value=str(registro_editar.get("OBSERVACIONES", ""))
                    )

                bx, by = st.columns([4, 1])
                with bx:
                    guardar_edicion_salida = st.form_submit_button(
                        "💾 ACTUALIZAR REGISTRO", use_container_width=True
                    )
                with by:
                    cancelar_edicion_salida = st.form_submit_button(
                        "✕ CANCELAR", type="primary", use_container_width=True
                    )

            if cancelar_edicion_salida:
                # Cierra el formulario de edición y deselecciona el registro en la tabla.
                st.session_state.salida_id_editar_inline = None
                st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                st.rerun()

            if guardar_edicion_salida:
                nuevo_consumo_cajas_nuevas = int(edit_cajas) - int(edit_reutilizadas)
                if not edit_solicitante.strip() or not edit_agencia:
                    st.error("Completa SOLICITANTE y AGENCIA antes de actualizar.")
                elif int(edit_cajas) <= 0 and int(edit_cintillos) <= 0:
                    st.error("Debes registrar al menos una cantidad mayor a cero de CAJAS o CINTILLOS.")
                elif nuevo_consumo_cajas_nuevas < 0:
                    st.error("Las cajas reutilizadas no pueden ser mayores que las cajas solicitadas.")
                else:
                    viejo_cajas = entero_seguro(registro_editar.get("CANTIDAD DE CAJAS SOLICITADAS", 0))
                    viejo_reutilizadas = entero_seguro(registro_editar.get("CAJAS REUTILIZADAS", 0))
                    viejo_cintillos = entero_seguro(registro_editar.get("CANTIDAD DE CINTILLOS", 0))
                    viejo_consumo_nuevas = viejo_cajas - viejo_reutilizadas

                    # Ajuste necesario para que el stock refleje exactamente el registro editado.
                    ajuste_nuevas = viejo_consumo_nuevas - nuevo_consumo_cajas_nuevas
                    ajuste_reutilizadas = viejo_reutilizadas - int(edit_reutilizadas)
                    ajuste_cintillos = viejo_cintillos - int(edit_cintillos)

                    stock_actual_edicion = obtener_stock_actual()
                    falta_nuevas = max(0, -ajuste_nuevas - stock_actual_edicion["CAJAS NUEVAS"])
                    falta_reutilizadas = max(0, -ajuste_reutilizadas - stock_actual_edicion["CAJAS REUTILIZADAS"])
                    falta_cintillos = max(0, -ajuste_cintillos - stock_actual_edicion["CINTILLOS"])

                    if falta_nuevas > 0 or falta_reutilizadas > 0 or falta_cintillos > 0:
                        st.error("❌ No se puede actualizar porque el cambio dejaría el stock en negativo.")
                    else:
                        datos_editados = {
                            "SOLICITANTE": edit_solicitante.strip(),
                            "CLIENTE": edit_cliente,
                            "AGENCIA": edit_agencia.strip(),
                            "FECHA DE SALIDA": edit_fecha.isoformat(),
                            "CANTIDAD DE CAJAS SOLICITADAS": int(edit_cajas),
                            "NRO. WORKORDER": edit_wo.strip(),
                            "CAJAS REUTILIZADAS": int(edit_reutilizadas),
                            "CANTIDAD DE CINTILLOS": int(edit_cintillos),
                            "OBSERVACIONES": edit_observaciones.strip(),
                        }
                        try:
                            actualizar_salida_cajas_supabase(datos_editados, int(salida_id_editar))
                            if ajuste_nuevas != 0 or ajuste_reutilizadas != 0 or ajuste_cintillos != 0:
                                guardar_movimiento_stock(
                                    edit_fecha,
                                    "AJUSTE",
                                    ajuste_nuevas,
                                    ajuste_reutilizadas,
                                    ajuste_cintillos,
                                    "",
                                    f"Ajuste por edición de salida ID {int(salida_id_editar)} · WO: {edit_wo.strip()}",
                                )
                            st.session_state.cajas = cargar_cajas_supabase()
                            st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
                            st.session_state.salida_id_editar_inline = None
                            st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                            st.success("✅ Registro actualizado correctamente y stock ajustado.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ No se pudo actualizar el registro: {e}")
    else:
        tabla_filtrada = cajas_df.copy()

    if tabla_filtrada.empty:
        if cajas_df.empty:
            st.info("No hay salidas registradas todavía.")
        else:
            st.info("No hay salidas que coincidan con los filtros seleccionados.")
    else:
        # ADMIN edita seleccionando la casilla ✏️ directamente al lado de cada registro,
        # igual que en la tabla principal de Solicitudes.
        if es_admin():
            tabla_cajas = tabla_filtrada.copy()
            tabla_cajas["✏️ EDITAR"] = False
            tabla_cajas["🗑️ ELIMINAR"] = False
            resultado_cajas = st.data_editor(
                tabla_cajas,
                use_container_width=True,
                hide_index=True,
                height=420,
                key=f"tabla_historial_salidas_{st.session_state.get('caja_tabla_version', 0)}",
                disabled=[c for c in tabla_cajas.columns if c not in ["✏️ EDITAR", "🗑️ ELIMINAR"]],
                column_config={
                    "_ID_CAJA_": None,
                    "FECHA DE SALIDA": st.column_config.DateColumn(
                        "FECHA DE SALIDA", format="DD/MM/YYYY", width=130
                    ),
                    "✏️ EDITAR": st.column_config.CheckboxColumn(
                        "EDIT.", default=False, width=60
                    ),
                    "🗑️ ELIMINAR": st.column_config.CheckboxColumn(
                        "ELIM.", default=False, width=65
                    ),
                },
            )

            # Editar directamente desde la tabla.
            editar_caja = resultado_cajas[resultado_cajas["✏️ EDITAR"] == True]
            if not editar_caja.empty:
                salida_seleccionada = int(editar_caja.iloc[0]["_ID_CAJA_"])
                if st.session_state.get("salida_id_editar_inline") != salida_seleccionada:
                    st.session_state.salida_id_editar_inline = salida_seleccionada
                    st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                    st.rerun()

            # Eliminar desde la misma tabla, con confirmación para evitar borrados accidentales.
            eliminar_caja = resultado_cajas[resultado_cajas["🗑️ ELIMINAR"] == True]
            st.session_state.setdefault("salida_id_eliminar_inline", None)
            if not eliminar_caja.empty:
                id_eliminar = int(eliminar_caja.iloc[0]["_ID_CAJA_"])
                if st.session_state.get("salida_id_eliminar_inline") != id_eliminar:
                    st.session_state.salida_id_eliminar_inline = id_eliminar
                    st.session_state.salida_id_editar_inline = None
                    st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                    st.rerun()

            id_eliminar = st.session_state.get("salida_id_eliminar_inline")
            if id_eliminar is not None:
                registro_eliminar = tabla_filtrada[
                    tabla_filtrada["_ID_CAJA_"].astype(int) == int(id_eliminar)
                ]
                if registro_eliminar.empty:
                    st.session_state.salida_id_eliminar_inline = None
                else:
                    reg_eliminar = registro_eliminar.iloc[0]
                    st.warning(
                        f"⚠️ ¿Deseas eliminar la salida de {reg_eliminar.get('SOLICITANTE', '')} "
                        f"del {convertir_fecha(reg_eliminar.get('FECHA DE SALIDA')).strftime('%d/%m/%Y') if convertir_fecha(reg_eliminar.get('FECHA DE SALIDA')) else reg_eliminar.get('FECHA DE SALIDA', '')}? "
                        "El stock consumido por este registro será devuelto automáticamente."
                    )
                    ce1, ce2 = st.columns([1, 1])
                    with ce1:
                        confirmar_eliminacion = st.button(
                            "🗑️ CONFIRMAR ELIMINACIÓN",
                            type="primary",
                            use_container_width=True,
                            key=f"confirmar_eliminar_caja_{id_eliminar}",
                        )
                    with ce2:
                        cancelar_eliminacion = st.button(
                            "✕ CANCELAR",
                            use_container_width=True,
                            key=f"cancelar_eliminar_caja_{id_eliminar}",
                        )

                    if cancelar_eliminacion:
                        st.session_state.salida_id_eliminar_inline = None
                        st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                        st.rerun()

                    if confirmar_eliminacion:
                        try:
                            cajas_eliminar = entero_seguro(reg_eliminar.get("CANTIDAD DE CAJAS SOLICITADAS", 0))
                            reutilizadas_eliminar = entero_seguro(reg_eliminar.get("CAJAS REUTILIZADAS", 0))
                            cintillos_eliminar = entero_seguro(reg_eliminar.get("CANTIDAD DE CINTILLOS", 0))
                            nuevas_eliminar = cajas_eliminar - reutilizadas_eliminar
                            fecha_eliminar = convertir_fecha(reg_eliminar.get("FECHA DE SALIDA")) or fecha_local_hoy()
                            wo_eliminar = str(reg_eliminar.get("NRO. WORKORDER", "")).strip()

                            # Primero elimina la salida y luego devuelve al stock lo que esa salida consumió.
                            eliminar_salida_cajas_supabase(int(id_eliminar))
                            if nuevas_eliminar or reutilizadas_eliminar or cintillos_eliminar:
                                guardar_movimiento_stock(
                                    fecha_eliminar,
                                    "DEVOLUCIÓN POR ELIMINACIÓN",
                                    nuevas_eliminar,
                                    reutilizadas_eliminar,
                                    cintillos_eliminar,
                                    "",
                                    f"Devolución por eliminación de salida ID {int(id_eliminar)} · WO: {wo_eliminar}",
                                )

                            st.session_state.cajas = cargar_cajas_supabase()
                            st.session_state.movimientos_stock = cargar_movimientos_stock_supabase()
                            st.session_state.salida_id_eliminar_inline = None
                            st.session_state.salida_id_editar_inline = None
                            st.session_state["caja_tabla_version"] = st.session_state.get("caja_tabla_version", 0) + 1
                            st.success("✅ Registro eliminado correctamente y stock devuelto.")
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ No se pudo eliminar el registro: {e}")
        else:
            tabla_cajas = tabla_filtrada.drop(columns=["_ID_CAJA_"], errors="ignore").copy()
            st.dataframe(
                tabla_cajas,
                use_container_width=True,
                hide_index=True,
                height=420,
                column_config={
                    "FECHA DE SALIDA": st.column_config.DateColumn(
                        "FECHA DE SALIDA", format="DD/MM/YYYY", width=130
                    )
                },
            )

    with st.expander("📊 VER HISTORIAL DE MOVIMIENTOS DE STOCK", expanded=False):
        mov = cargar_movimientos_stock_supabase()
        if mov.empty:
            st.info("No hay movimientos registrados todavía.")
        else:
            columnas_mov = [c for c in ["fecha", "tipo_movimiento", "cajas_nuevas", "cajas_reutilizadas", "cintillos", "proveedor", "observaciones", "usuario"] if c in mov.columns]

            # Mostrar fecha en formato DD/MM/YY sin modificar el dato original de Supabase.
            mov_mostrar = mov[columnas_mov].copy()
            if "fecha" in mov_mostrar.columns:
                mov_mostrar["fecha"] = pd.to_datetime(
                    mov_mostrar["fecha"], errors="coerce"
                ).dt.strftime("%d/%m/%y").fillna("")

            st.dataframe(
                mov_mostrar,
                use_container_width=True,
                hide_index=True,
                height=350
            )

    st.markdown('<div class="footer"><strong>JC Control de Solicitudes — Almacén</strong><br>©JuanCarlosRamos - 2026 — Todos los derechos reservados</div>', unsafe_allow_html=True)
    st.stop()

# =========================================================
# PAGINA PROGRAMACION DE RUTAS
# =========================================================
if st.session_state.page == "rutas":
    st.subheader("🗓️ PROGRAMACIÓN DE RUTAS")
    st.caption("Registro y consulta de programación directamente desde Supabase PostgreSQL.")

    if st.button("⬅ VOLVER A SOLICITUDES"):
        liberar_bloqueo_edicion()
        st.session_state.page = "solicitudes"
        st.rerun()

    if puede_modificar():
        bloqueo = leer_bloqueo_edicion()
        if bloqueo and bloqueo.get("owner_id") != st.session_state.session_id:
            st.warning(f"🔒 {bloqueo.get('usuario', 'Otro usuario')} está realizando una operación de edición. Espera para registrar una ruta.")
        else:
            if not st.session_state.ruta_activo:
                if st.button("🔐 INICIAR REGISTRO DE RUTA"):
                    ok, info = adquirir_bloqueo_edicion("nueva programación de ruta")
                    if ok:
                        st.session_state.ruta_activo = True
                        st.session_state.ruta_form_version += 1
                        st.rerun()
                    else:
                        st.warning(f"🔒 {info.get('usuario', 'Otro usuario')} está utilizando el sistema.")

        if st.session_state.ruta_activo:
            renovar_bloqueo_edicion()
            st.success("🔐 Registro de ruta activo. La edición está reservada para este usuario.")
            formulario_nueva_ruta()
    else:
        st.info("👁️ MODO INVITADO — Puedes consultar las programaciones, pero no registrar ni modificar datos.")

    st.markdown("### 📋 Programaciones registradas")
    rutas = st.session_state.rutas
    if rutas is None or rutas.empty:
        st.info("No hay programaciones registradas todavía.")
    else:
        f1, f2 = st.columns([1, 2])
        with f1:
            regional_filtro = st.selectbox("Regional", ["TODOS"] + REGIONALES[1:], key="ruta_regional")
        with f2:
            buscar_ruta = st.text_input("Buscar", placeholder="Centro de acopio, agencia o regional...", key="buscar_ruta")

        datos = rutas.copy()
        if regional_filtro != "TODOS":
            datos = datos[datos["REGIONAL"].map(normalizar).eq(normalizar(regional_filtro))]
        if buscar_ruta:
            texto = normalizar(buscar_ruta)
            mask = datos.astype(str).apply(lambda col: col.map(lambda x: texto in normalizar(x))).any(axis=1)
            datos = datos[mask]

        st.caption(f"📋 {len(datos)} de {len(rutas)} programaciones encontradas")
        tabla_rutas = datos.drop(columns=["_ID_RUTA_"], errors="ignore").copy()

        # Formato de fechas para que la tabla sea más compacta y uniforme.
        for col in [
            "FECHA LIMITE DE INGRESO (SE1)",
            "FECHA LIMITE DE INGRESO (SR1)",
            "FECHA DE RECOJO",
        ]:
            if col in tabla_rutas.columns:
                tabla_rutas[col] = pd.to_datetime(
                    tabla_rutas[col], errors="coerce"
                )

        st.dataframe(
            tabla_rutas,
            use_container_width=True,
            hide_index=True,
            height=500,
            column_config={
                "REGIONAL": st.column_config.TextColumn(
                    "REGIONAL", width=120
                ),
                "CENTRO DE ACOPIO": st.column_config.TextColumn(
                    "CENTRO DE ACOPIO", width=190
                ),
                "AGENCIAS": st.column_config.TextColumn(
                    "AGENCIAS", width=250
                ),
                "FECHA LIMITE DE INGRESO (SE1)": st.column_config.DateColumn(
                    "FECHA LÍMITE SE1", format="DD/MM/YYYY", width=150
                ),
                "FECHA LIMITE DE INGRESO (SR1)": st.column_config.DateColumn(
                    "FECHA LÍMITE SR1", format="DD/MM/YYYY", width=150
                ),
                "FECHA DE RECOJO": st.column_config.DateColumn(
                    "FECHA DE RECOJO", format="DD/MM/YYYY", width=145
                ),
            },
        )
        _, col_descarga_ruta, _ = st.columns([1, 2, 1])
        with col_descarga_ruta:
            st.download_button(
                "📥 DESCARGAR PROGRAMACIÓN",
                data=excel_bytes(st.session_state.rows, datos),
                file_name=f"PROGRAMACION_RUTAS_{date.today().isoformat()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

    st.markdown('<div class="footer"><strong>JC Control de Solicitudes — Almacén</strong><br>©JuanCarlosRamos - 2026 — Todos los derechos reservados</div>', unsafe_allow_html=True)
    st.stop()


# =========================================================
# PAGINA SOLICITUDES
# =========================================================
df = st.session_state.rows
st.markdown("### 📁 Base de datos compartida")
st.caption("🗄️ Supabase PostgreSQL · Datos compartidos entre todos los usuarios.")

bloqueo_actual = leer_bloqueo_edicion()
if bloqueo_actual and bloqueo_actual.get("owner_id") != st.session_state.session_id:
    st.warning(f"🔒 BASE DE DATOS EN USO — {bloqueo_actual.get('usuario', 'otro usuario')} está registrando o editando. Puedes consultar, pero espera para modificar.")

# =========================================================
# NUEVA SOLICITUD
# =========================================================
if puede_modificar() and st.session_state.editing is None and not st.session_state.registro_activo:
    if st.button("🔐 INICIAR NUEVA SOLICITUD"):
        ok, info = adquirir_bloqueo_edicion("nueva solicitud")
        if ok:
            st.session_state.registro_activo = True
            st.session_state.form_version += 1
            st.rerun()
        else:
            st.warning(f"🔒 {info.get('usuario', 'Otro usuario')} está utilizando el sistema.")

# =========================================================
# FORMULARIO SOLICITUD
# =========================================================
if puede_modificar() and (st.session_state.registro_activo or st.session_state.editing is not None):
    renovar_bloqueo_edicion()
    editing = st.session_state.editing
    if editing is not None and editing in df.index:
        r = df.loc[editing]
        valores = {"cliente":r["CLIENTE"],"numero":r["NRO SOLICITUD - WO"],"tipo":r["TIPO DE SOLICITUD"],"centro_costo":r.get("CENTRO DE COSTO",""),"prioridad":r["PRIORIDAD"],"cantidad_items":r["CANT - ITEMS"],"estado":r["ESTADO DE SOLICITUD"],"direccion":r["DIRECCIÓN"],"fecha":r["FECHA DE INGRESO"]}
        titulo_form = "✏️ Editar solicitud"
    else:
        valores = {"cliente":CAT_CLIENTES[0],"numero":"","tipo":CAT_TIPOS[0],"centro_costo":"","prioridad":PRIORIDADES[0],"cantidad_items":0,"estado":CAT_ESTADOS[0],"direccion":CAT_DIRECCIONES[0],"fecha":str(date.today())}
        titulo_form = "📥 Registrar nueva solicitud"

    with st.form(f"solicitud_form_{st.session_state.form_version}"):
        st.markdown(f"**{titulo_form}**")
        a,b = st.columns(2)
        with a:
            cliente = st.selectbox("Cliente", CAT_CLIENTES, index=CAT_CLIENTES.index(valores["cliente"]) if valores["cliente"] in CAT_CLIENTES else 0)
            numero = st.text_input("Nro. Solicitud / WO", value=str(valores["numero"] or ""), placeholder="Ejemplo: SE2-26-17658")
            tipo = st.selectbox("Tipo de Solicitud", CAT_TIPOS, index=CAT_TIPOS.index(valores["tipo"]) if valores["tipo"] in CAT_TIPOS else 0)
            centro_actual = str(valores.get("centro_costo", "") or "").strip()
            opciones_centro = CAT_CENTROS_COSTO.copy()
            if centro_actual and centro_actual not in opciones_centro:
                opciones_centro.append(centro_actual)
            centro_costo = st.selectbox(
                "Centro de Costo",
                opciones_centro,
                index=opciones_centro.index(centro_actual) if centro_actual in opciones_centro else 0,
            )
            prioridad = st.selectbox("Prioridad", PRIORIDADES, index=PRIORIDADES.index(valores["prioridad"]) if valores["prioridad"] in PRIORIDADES else 0)
        with b:
            try: cantidad_default = max(0, int(float(str(valores.get("cantidad_items",0)).replace(",","."))))
            except Exception: cantidad_default = 0
            cantidad_items = st.number_input("Cantidad de ítems", min_value=0, step=1, value=cantidad_default, format="%d")
            estado = st.selectbox("Estado de Solicitud", CAT_ESTADOS, index=CAT_ESTADOS.index(valores["estado"]) if valores["estado"] in CAT_ESTADOS else 0)
            direccion = st.selectbox("Dirección", CAT_DIRECCIONES, index=CAT_DIRECCIONES.index(valores["direccion"]) if valores["direccion"] in CAT_DIRECCIONES else 0)
            fecha_default = convertir_fecha(valores.get("fecha")) or fecha_local_hoy()
            fecha = st.date_input("📅 Fecha de ingreso", value=fecha_default, format="DD/MM/YYYY")
        x,y = st.columns([4,1])
        with x: guardar = st.form_submit_button("💾 ACTUALIZAR REGISTRO" if editing is not None else "📥 AGREGAR REGISTRO", use_container_width=True)
        with y: cancelar = st.form_submit_button("✕ CANCELAR", type="primary", use_container_width=True)

    if cancelar:
        liberar_bloqueo_edicion()
        st.session_state.editing = None
        st.session_state.editing_key = ""
        st.session_state.form_version += 1
        st.rerun()

    if guardar:
        if editing is not None and not puede_modificar():
            st.error("⛔ Tu usuario no tiene permiso para editar solicitudes.")
            st.stop()
        if editing is None and not puede_modificar():
            st.error("⛔ Tu usuario no tiene permiso para registrar solicitudes.")
            st.stop()
        numero_norm = normalizar(numero)
        if not numero_norm or cliente == CAT_CLIENTES[0] or tipo == CAT_TIPOS[0] or prioridad == PRIORIDADES[0] or estado == CAT_ESTADOS[0] or direccion == CAT_DIRECCIONES[0]:
            st.error("Completa todos los campos obligatorios antes de guardar.")
        else:
            try:
                latest = cargar_solicitudes_supabase()
                nuevo = {"CLIENTE":cliente,"NRO SOLICITUD - WO":numero.strip(),"TIPO DE SOLICITUD":tipo,"CENTRO DE COSTO":centro_costo.strip(),"PRIORIDAD":prioridad,"CANT - ITEMS":int(cantidad_items),"ESTADO DE SOLICITUD":estado,"DIRECCIÓN":direccion,"FECHA DE INGRESO":fecha.isoformat()}
                repetidos = latest[latest["NRO SOLICITUD - WO"].map(normalizar).eq(numero_norm)]
                if editing is not None:
                    registro_id = int(df.loc[editing,"_ID_"])
                    if not repetidos.empty and int(repetidos.iloc[0]["_ID_"]) != registro_id:
                        st.error("Ya existe otra solicitud con ese Nro. Solicitud / WO.")
                    else:
                        guardar_solicitud_supabase(nuevo, registro_id)
                        st.session_state.rows = cargar_solicitudes_supabase()
                        liberar_bloqueo_edicion()
                        st.session_state.editing = None
                        st.session_state.editing_key = ""
                        st.success("✅ Registro actualizado correctamente en Supabase.")
                        st.rerun()
                else:
                    if not repetidos.empty:
                        st.error("Ya existe una solicitud con ese Nro. Solicitud / WO.")
                    else:
                        guardar_solicitud_supabase(nuevo)
                        st.session_state.rows = cargar_solicitudes_supabase()
                        liberar_bloqueo_edicion()
                        st.success("✅ Nueva solicitud agregada correctamente en Supabase.")
                        st.rerun()
            except Exception as e:
                liberar_bloqueo_edicion()
                st.error(f"❌ No se pudo guardar en Supabase: {e}")

# =========================================================
# FILTROS
# =========================================================
st.markdown("### Resultados")
buscar_col, cliente_col, estado_col, fecha_col = st.columns([3,1.4,1.4,1.2])
with buscar_col:
    buscar = st.text_input("Buscar solicitud...", placeholder="Cliente, WO, dirección, tipo...", key="buscar_solicitud_live")
with cliente_col:
    filtro_cliente = st.selectbox("Cliente", ["TODOS"] + CAT_CLIENTES[1:], key="filtro_cliente")
with estado_col:
    filtro_estado = st.selectbox("Estado", ["TODOS"] + CAT_ESTADOS[1:], key="filtro_estado_solicitud")
with fecha_col:
    filtro_fecha = st.date_input("📅 Fecha", value=None, key="filtro_fecha_solicitud", format="DD/MM/YYYY")

# =========================================================
# BÚSQUEDA MASIVA DE SOLICITUDES
# =========================================================
with st.expander("🔎 BÚSQUEDA MASIVA — pegar solicitudes o cargar Excel/TXT", expanded=False):
    st.caption("Puedes pegar varios NRO SOLICITUD / WO, uno por línea, o cargar un archivo .xlsx, .csv o .txt. Se buscarán coincidencias exactas en la columna NRO SOLICITUD - WO.")
    bm1, bm2 = st.columns([3, 2])
    with bm1:
        texto_masivo = st.text_area(
            "Lista de solicitudes",
            placeholder="SE1-26-00001\nSE2-26-00002\nSR1-26-00003",
            height=110,
            key=f"texto_busqueda_masiva_{st.session_state.busqueda_masiva_version}",
        )
    with bm2:
        archivo_masivo = st.file_uploader(
            "Cargar archivo",
            type=["xlsx", "csv", "txt"],
            key=f"archivo_busqueda_masiva_{st.session_state.busqueda_masiva_version}",
            help="El archivo puede contener una o varias columnas; el sistema comparará sus valores con los NRO SOLICITUD - WO existentes.",
        )
        procesar_masivo = st.button("🔍 BUSCAR SOLICITUDES", use_container_width=True, key="procesar_busqueda_masiva")
        limpiar_masivo = st.button("✕ LIMPIAR BÚSQUEDA MASIVA", use_container_width=True, key="limpiar_busqueda_masiva")

    if limpiar_masivo:
        st.session_state.busqueda_masiva_ids = []
        st.session_state.busqueda_masiva_nombre = ""
        # Cambiar las keys hace que Streamlit cree los widgets nuevamente,
        # borrando el texto y el archivo cargado. El expander vuelve a su
        # estado expanded=False y queda automáticamente replegado.
        st.session_state.busqueda_masiva_version += 1
        st.session_state.tabla_version += 1
        st.rerun()

    if procesar_masivo:
        try:
            ids_texto = extraer_solicitudes_de_texto(texto_masivo)
            ids_archivo = extraer_solicitudes_de_archivo(archivo_masivo) if archivo_masivo is not None else []
            candidatos = []
            for valor in ids_texto + ids_archivo:
                valor_norm = normalizar(valor)
                if valor_norm and valor_norm not in candidatos:
                    candidatos.append(valor_norm)

            ids_existentes = set(df["NRO SOLICITUD - WO"].map(normalizar).tolist()) if "NRO SOLICITUD - WO" in df.columns else set()
            encontrados = [x for x in candidatos if x in ids_existentes]
            no_encontrados = [x for x in candidatos if x not in ids_existentes]
            st.session_state.busqueda_masiva_ids = encontrados
            st.session_state.busqueda_masiva_nombre = getattr(archivo_masivo, "name", "lista pegada") if archivo_masivo is not None else "lista pegada"

            if encontrados:
                st.success(f"✅ {len(encontrados)} solicitud(es) encontrada(s) de {len(candidatos)} buscada(s).")
            elif candidatos:
                st.warning("⚠️ No se encontraron solicitudes que coincidan exactamente con la base de datos.")
            else:
                st.warning("⚠️ Ingresa al menos un NRO SOLICITUD / WO o carga un archivo.")

            if no_encontrados and encontrados:
                st.caption(f"No encontradas: {len(no_encontrados)}")
        except Exception as e:
            st.error(f"❌ No se pudo procesar la búsqueda masiva: {e}")

    if st.session_state.busqueda_masiva_ids:
        st.info(f"🔎 Búsqueda masiva activa: {len(st.session_state.busqueda_masiva_ids)} solicitud(es). Fuente: {st.session_state.busqueda_masiva_nombre or 'lista pegada'}")

vista = df.copy()
if st.session_state.busqueda_masiva_ids:
    ids_masivos = set(st.session_state.busqueda_masiva_ids)
    vista = vista[vista["NRO SOLICITUD - WO"].map(normalizar).isin(ids_masivos)]
if buscar:
    texto = normalizar(buscar)
    vista = vista[vista.astype(str).apply(lambda col: col.map(lambda x: texto in normalizar(x))).any(axis=1)]
if filtro_cliente != "TODOS":
    vista = vista[vista["CLIENTE"].map(normalizar).eq(normalizar(filtro_cliente))]
if filtro_estado != "TODOS":
    vista = vista[vista["ESTADO DE SOLICITUD"].map(normalizar).eq(normalizar(filtro_estado))]
if filtro_fecha:
    fechas = vista["FECHA DE INGRESO"].apply(convertir_fecha)
    vista = vista[fechas.apply(lambda f: f == filtro_fecha)]

# =========================================================
# ALERTAS DE ATENCIÓN POR ANTIGÜEDAD
# =========================================================
# Las solicitudes ENVIADAS, ENTREGADAS o ANULADAS no generan alertas.
# 1 día de antigüedad  -> advertencia
# 2 o más días          -> atención urgente
estados_sin_alerta = {"ENVIADO", "ENTREGADO", "ANULADO"}

alertas = df.copy()

# Convertir cada fecha con la misma lógica usada por el sistema:
# - ISO de Supabase: YYYY-MM-DD
# - Fecha mostrada/guardada como DD/MM/YYYY
# Esto evita que las solicitudes antiguas se interpreten con mes y día invertidos.
alertas["_FECHA_ALERTA"] = alertas["FECHA DE INGRESO"].apply(convertir_fecha)

hoy_app = fecha_local_hoy()
alertas["_DIAS_ALERTA"] = alertas["_FECHA_ALERTA"].apply(
    lambda f: (hoy_app - f).days if f is not None else None
)

alertas = alertas[
    (~alertas["ESTADO DE SOLICITUD"].astype(str).str.strip().str.upper().isin(estados_sin_alerta))
    & alertas["_FECHA_ALERTA"].notna()
    & (alertas["_DIAS_ALERTA"] >= 1)
]

if not alertas.empty:
    urgentes = alertas[alertas["_DIAS_ALERTA"] >= 2]
    advertencias = alertas[alertas["_DIAS_ALERTA"] == 1]

    st.markdown("### 🔔 Alertas de atención")
    met1, met2, met3 = st.columns(3)
    with met1:
        st.metric("🚨 Atención urgente (2+ días)", len(urgentes))
    with met2:
        st.metric("⚠️ Atención (1 día)", len(advertencias))
    with met3:
        st.metric("📌 Total pendientes", len(alertas))

    # Botón desplegable para no ocupar espacio con las tablas de alertas.
    with st.expander("🔽 Ver tabla de alertas de atención", expanded=False):
        if not urgentes.empty:
            st.error(
                f"🚨 Hay **{len(urgentes)} solicitud(es)** con 2 o más días de antigüedad que requieren atención urgente."
            )
            detalle_urg = urgentes[["NRO SOLICITUD - WO", "CLIENTE", "FECHA DE INGRESO", "ESTADO DE SOLICITUD", "_DIAS_ALERTA"]].copy()
            detalle_urg["FECHA DE INGRESO"] = detalle_urg["FECHA DE INGRESO"].apply(
                lambda f: convertir_fecha(f).strftime("%d/%m/%Y") if convertir_fecha(f) else ""
            )
            detalle_urg["ATENCIÓN"] = detalle_urg["_DIAS_ALERTA"].map(lambda x: f"🚨 {int(x)} días")
            detalle_urg = detalle_urg.drop(columns=["_DIAS_ALERTA"])
            st.dataframe(detalle_urg, use_container_width=True, hide_index=True)

        if not advertencias.empty:
            st.warning(
                f"⚠️ Hay **{len(advertencias)} solicitud(es)** que llevan 1 día desde su registro."
            )
            detalle_adv = advertencias[["NRO SOLICITUD - WO", "CLIENTE", "FECHA DE INGRESO", "ESTADO DE SOLICITUD"]].copy()
            detalle_adv["FECHA DE INGRESO"] = detalle_adv["FECHA DE INGRESO"].apply(
                lambda f: convertir_fecha(f).strftime("%d/%m/%Y") if convertir_fecha(f) else ""
            )
            detalle_adv["ATENCIÓN"] = "⚠️ 1 día"
            st.dataframe(detalle_adv, use_container_width=True, hide_index=True)
else:
    st.success("✅ No hay solicitudes pendientes con más de 1 día de antigüedad que requieran atención.")

st.caption(f"{len(vista)} de {len(df)} registros")

if len(vista):
    tabla = vista[COLS].copy()
    tabla["_INDICE_REAL_"] = vista.index
    tabla["☑️ SELECCIONAR"] = False
    tabla["✏️ EDITAR"] = False
    tabla["🗑️ ELIMINAR"] = False

    # =====================================================
    # COLORES VISUALES PARA ESTADO Y PRIORIDAD
    # =====================================================
    # Se muestran como indicadores de color dentro de la
    # tabla sin modificar los valores originales de Supabase.
    estado_colores = {
        "POR EXTRAER": "🟡 POR EXTRAER",
        "POR REASIGNAR": "🟣 POR REASIGNAR",
        "POR ENVIAR": "🟠 POR ENVIAR",
        "ENTREGADO": "🟢 ENTREGADO",
        "ENVIADO": "🔵 ENVIADO",
        "ANULADO": "🔴 ANULADO",
        "POR ETIQUETAR": "🟡 POR ETIQUETAR",
    }

    prioridad_colores = {
        "RUSH": "🔴 RUSH",
        "TURNO SIGUIENTE": "🟠 TURNO SIGUIENTE",
        "NORMAL": "🟢 NORMAL",
    }

    tabla["ESTADO DE SOLICITUD"] = (
        tabla["ESTADO DE SOLICITUD"]
        .map(lambda x: estado_colores.get(str(x).strip(), str(x)))
    )

    tabla["PRIORIDAD"] = (
        tabla["PRIORIDAD"]
        .map(lambda x: prioridad_colores.get(str(x).strip(), str(x)))
    )
    resultado = st.data_editor(
        tabla,
        use_container_width=True,
        hide_index=True,
        height=470,
        key=f"tabla_solicitudes_{st.session_state.tabla_version}",
        disabled=COLS + ["_INDICE_REAL_"] + (
            [] if puede_modificar() else ["☑️ SELECCIONAR", "✏️ EDITAR", "🗑️ ELIMINAR"]
        ),
        column_config={
            "_INDICE_REAL_": None,
            "☑️ SELECCIONAR": st.column_config.CheckboxColumn(
                "SELEC.", default=False, width=65, disabled=not puede_modificar()
            ),

            # Anchos fijos para que todas las columnas entren de forma
            # equilibrada en la pantalla.
            "CLIENTE": st.column_config.TextColumn("CLIENTE", width=130),
            "NRO SOLICITUD - WO": st.column_config.TextColumn("N° SOLICITUD / WO", width=160),
            "TIPO DE SOLICITUD": st.column_config.TextColumn("TIPO", width=125),
            "CENTRO DE COSTO": st.column_config.TextColumn("CENTRO DE COSTO", width=125),
            "PRIORIDAD": st.column_config.TextColumn("PRIORIDAD", width=160),
            "CANT - ITEMS": st.column_config.NumberColumn(
                "ÍTEMS", min_value=0, step=1, format="%d", width=70
            ),
            "ESTADO DE SOLICITUD": st.column_config.TextColumn("ESTADO", width=130),
            "DIRECCIÓN": st.column_config.TextColumn("DIRECCIÓN", width=150),
            "FECHA DE INGRESO": st.column_config.DateColumn(
                "FECHA", format="DD/MM/YYYY", width=105
            ),
            "✏️ EDITAR": st.column_config.CheckboxColumn("EDIT.", default=False, width=60),
            "🗑️ ELIMINAR": st.column_config.CheckboxColumn(
                "ELIM.", default=False, width=60, disabled=not es_admin()
            ),
        },
    )

    # =====================================================
    # EDICIÓN MASIVA DE SOLICITUDES
    # =====================================================
    seleccion_solicitudes = resultado[
        resultado["☑️ SELECCIONAR"] == True
    ].copy()

    if puede_modificar() and not seleccion_solicitudes.empty:
        indices_masivos = [
            int(x) for x in seleccion_solicitudes["_INDICE_REAL_"].tolist()
        ]

        with st.expander(
            f"🛠️ EDITAR {len(indices_masivos)} SOLICITUD(ES) SELECCIONADA(S)",
            expanded=True,
        ):
            st.info(
                "Selecciona varias solicitudes en **SELEC.** y aplica el mismo "
                "cambio a todas. Los demás datos de cada solicitud permanecerán "
                "sin modificaciones."
            )

            campo_masivo = st.selectbox(
                "Campo que deseas modificar",
                [
                    "ESTADO DE SOLICITUD",
                    "PRIORIDAD",
                    "CENTRO DE COSTO",
                    "CLIENTE",
                    "TIPO DE SOLICITUD",
                    "DIRECCIÓN",
                    "CANT - ITEMS",
                ],
                key=f"campo_edicion_masiva_{st.session_state.tabla_version}",
            )

            valores_masivos = {
                "ESTADO DE SOLICITUD": CAT_ESTADOS[1:],
                "PRIORIDAD": PRIORIDADES[1:],
                "CENTRO DE COSTO": CAT_CENTROS_COSTO[1:],
                "CLIENTE": CAT_CLIENTES[1:],
                "TIPO DE SOLICITUD": CAT_TIPOS[1:],
                "DIRECCIÓN": CAT_DIRECCIONES[1:],
            }

            if campo_masivo == "CANT - ITEMS":
                valor_masivo = st.number_input(
                    "Nuevo valor",
                    min_value=0,
                    step=1,
                    value=0,
                    key=f"valor_items_masivo_{st.session_state.tabla_version}",
                )
            else:
                valor_masivo = st.selectbox(
                    "Nuevo valor",
                    valores_masivos[campo_masivo],
                    key=f"valor_edicion_masiva_{st.session_state.tabla_version}",
                )

            bm1, bm2 = st.columns([3, 1])
            with bm1:
                aplicar_masivo_solicitudes = st.button(
                    "💾 APLICAR CAMBIOS A LAS SOLICITUDES SELECCIONADAS",
                    type="primary",
                    use_container_width=True,
                    key=f"aplicar_edicion_masiva_solicitudes_{st.session_state.tabla_version}",
                )
            with bm2:
                cancelar_masivo_solicitudes = st.button(
                    "✕ CANCELAR",
                    use_container_width=True,
                    key=f"cancelar_edicion_masiva_solicitudes_{st.session_state.tabla_version}",
                )

            if cancelar_masivo_solicitudes:
                st.session_state.tabla_version += 1
                st.rerun()

            if aplicar_masivo_solicitudes:
                ok, info = adquirir_bloqueo_edicion(
                    "edición múltiple de solicitudes"
                )

                if not ok:
                    st.warning(
                        f"🔒 {info.get('usuario', 'Otro usuario')} está utilizando el sistema."
                    )
                else:
                    actualizados = 0
                    errores_masivos = []

                    for indice_masivo in indices_masivos:
                        try:
                            reg = df.loc[indice_masivo]

                            # Se toman TODOS los datos actuales y solo se
                            # reemplaza el campo seleccionado.
                            datos_m = {
                                "CLIENTE": str(reg.get("CLIENTE", "") or ""),
                                "NRO SOLICITUD - WO": str(
                                    reg.get("NRO SOLICITUD - WO", "") or ""
                                ),
                                "TIPO DE SOLICITUD": str(
                                    reg.get("TIPO DE SOLICITUD", "") or ""
                                ),
                                "CENTRO DE COSTO": str(
                                    reg.get("CENTRO DE COSTO", "") or ""
                                ),
                                "PRIORIDAD": str(reg.get("PRIORIDAD", "") or ""),
                                "CANT - ITEMS": int(
                                    pd.to_numeric(
                                        reg.get("CANT - ITEMS", 0), errors="coerce"
                                    )
                                    if pd.notna(
                                        pd.to_numeric(
                                            reg.get("CANT - ITEMS", 0),
                                            errors="coerce",
                                        )
                                    )
                                    else 0
                                ),
                                "ESTADO DE SOLICITUD": str(
                                    reg.get("ESTADO DE SOLICITUD", "") or ""
                                ),
                                "DIRECCIÓN": str(reg.get("DIRECCIÓN", "") or ""),
                                "FECHA DE INGRESO": (
                                    convertir_fecha(reg.get("FECHA DE INGRESO"))
                                    or fecha_local_hoy()
                                ).isoformat(),
                            }

                            if campo_masivo == "CANT - ITEMS":
                                datos_m["CANT - ITEMS"] = int(valor_masivo)
                            else:
                                datos_m[campo_masivo] = valor_masivo

                            guardar_solicitud_supabase(
                                datos_m,
                                int(reg.get("_ID_")),
                            )
                            actualizados += 1

                        except Exception as e:
                            errores_masivos.append(
                                f"ID {reg.get('_ID_', '?')}: {e}"
                            )

                    liberar_bloqueo_edicion()
                    st.session_state.rows = cargar_solicitudes_supabase()
                    st.session_state.tabla_version += 1

                    if actualizados:
                        st.success(
                            f"✅ Se actualizaron {actualizados} solicitud(es) correctamente."
                        )
                    if errores_masivos:
                        st.error(
                            "❌ Algunas solicitudes no pudieron actualizarse: "
                            + " | ".join(errores_masivos)
                        )

                    st.rerun()

    editar = resultado[resultado["✏️ EDITAR"] == True]
    if not puede_modificar():
        editar = resultado.iloc[0:0]
    if not editar.empty:
        indice_real = editar.iloc[0]["_INDICE_REAL_"]
        ok, info = adquirir_bloqueo_edicion("edición de solicitud")
        if ok:
            st.session_state.editing = int(indice_real)
            st.session_state.editing_key = normalizar(df.loc[int(indice_real), "NRO SOLICITUD - WO"])
            st.session_state.form_version += 1
            st.session_state.tabla_version += 1
            st.rerun()
        else:
            st.warning(f"🔒 {info.get('usuario', 'Otro usuario')} está editando.")

    eliminar = resultado[resultado["🗑️ ELIMINAR"] == True]
    if not es_admin():
        eliminar = resultado.iloc[0:0]
    if not eliminar.empty:
        ok, info = adquirir_bloqueo_edicion("eliminación de solicitud")
        if not ok:
            st.warning(f"🔒 {info.get('usuario', 'Otro usuario')} está utilizando el sistema.")
        else:
            ids = [int(df.loc[int(i), "_ID_"]) for i in eliminar["_INDICE_REAL_"].tolist()]
            try:
                eliminar_solicitudes_supabase(ids)
                st.session_state.rows = cargar_solicitudes_supabase()
                liberar_bloqueo_edicion()
                st.success("✅ Registro(s) eliminado(s) correctamente de Supabase.")
                st.session_state.tabla_version += 1
                st.rerun()
            except Exception as e:
                liberar_bloqueo_edicion()
                st.error(f"❌ No se pudo eliminar: {e}")
else:
    st.info("No hay registros que coincidan con los filtros.")

# =========================================================
# ARCHIVO MENSUAL Y LIMPIEZA DE BASE DE DATOS
# Solo ADMIN puede eliminar registros históricos.
# Primero se descarga el mes y luego se solicita confirmación
# explícita antes de borrar esos registros de Supabase.
# =========================================================
if es_admin():
    with st.expander("📦 ARCHIVAR Y LIMPIAR SOLICITUDES POR MES", expanded=False):
        st.warning(
            "⚠️ Esta opción descarga los registros de un mes y permite eliminarlos de Supabase. "
            "Se recomienda conservar el archivo Excel como respaldo antes de eliminar."
        )

        hoy = fecha_local_hoy()
        anos_disponibles = sorted(
            {d.year for d in pd.to_datetime(df["FECHA DE INGRESO"], errors="coerce").dropna()},
            reverse=True,
        )
        if hoy.year not in anos_disponibles:
            anos_disponibles.insert(0, hoy.year)
        meses = {
            1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL",
            5: "MAYO", 6: "JUNIO", 7: "JULIO", 8: "AGOSTO",
            9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE"
        }

        ac1, ac2 = st.columns(2)
        with ac1:
            anio_archivo = st.selectbox(
                "Año", anos_disponibles, key="anio_archivo_mensual"
            )
        with ac2:
            mes_archivo = st.selectbox(
                "Mes", list(meses.keys()),
                format_func=lambda m: meses[m],
                key="mes_archivo_mensual"
            )

        import calendar
        ultimo_dia = calendar.monthrange(anio_archivo, mes_archivo)[1]
        fecha_inicio_archivo = date(anio_archivo, mes_archivo, 1)
        fecha_fin_archivo = date(anio_archivo, mes_archivo, ultimo_dia)
        fecha_fin_exclusiva = date(anio_archivo + (mes_archivo == 12), 1 if mes_archivo == 12 else mes_archivo + 1, 1)

        # La confirmación de borrado pertenece exclusivamente al mes seleccionado.
        # Si el usuario cambia de año/mes, se limpia para evitar borrar otro mes por error.
        _mes_seleccionado = (anio_archivo, mes_archivo)
        _mes_previo = st.session_state.get("_mes_confirmacion_borrado")
        if _mes_previo != _mes_seleccionado:
            st.session_state.pop("confirmar_borrado_mensual", None)
            st.session_state["_mes_confirmacion_borrado"] = _mes_seleccionado

        # Después de un borrado exitoso se solicita nuevamente la confirmación.
        if st.session_state.pop("_reset_confirmar_borrado_mensual", False):
            st.session_state.pop("confirmar_borrado_mensual", None)

        try:
            registros_mes = cargar_solicitudes_mes_supabase(
                fecha_inicio_archivo, fecha_fin_exclusiva
            )
            cantidad_mes = len(registros_mes)

            st.info(
                f"📅 {meses[mes_archivo]} {anio_archivo}: **{cantidad_mes} registro(s)** encontrados "
                f"({fecha_inicio_archivo.strftime('%d/%m/%Y')} al {fecha_fin_archivo.strftime('%d/%m/%Y')})."
            )

            if cantidad_mes > 0:
                archivo_mes = excel_bytes_mensual(registros_mes[COLS])
                st.info(
                    "📊 El Excel incluirá el detalle completo con **NRO SOLICITUD - WO**, "
                    "además de resúmenes por estado, prioridad, SE1/SE2/SE3/SR1/SR2 y un resumen general."
                )

                _, col_descarga_mes, _ = st.columns([1, 2, 1])
                with col_descarga_mes:
                    st.download_button(
                        "📥 DESCARGAR ESTE MES EN EXCEL",
                        data=archivo_mes,
                        file_name=f"SOLICITUDES_{anio_archivo}_{mes_archivo:02d}_{meses[mes_archivo]}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="descargar_archivo_mensual",
                    )

                confirmar_borrado = st.checkbox(
                    "Confirmo que ya descargué y guardé el archivo de este mes y deseo eliminar estos registros de Supabase.",
                    key="confirmar_borrado_mensual",
                )

                if st.button(
                    f"🗑️ ELIMINAR {cantidad_mes} REGISTRO(S) DE {meses[mes_archivo]} {anio_archivo}",
                    type="secondary",
                    use_container_width=True,
                    disabled=not confirmar_borrado,
                    key="eliminar_mes_supabase",
                ):
                    ids_mes = [int(x) for x in registros_mes["_ID_"].tolist() if pd.notna(x)]
                    try:
                        eliminar_solicitudes_supabase(ids_mes)
                        st.session_state.rows = cargar_solicitudes_supabase()
                        st.session_state["_reset_confirmar_borrado_mensual"] = True
                        st.success(
                            f"✅ Se eliminaron {len(ids_mes)} registro(s) de {meses[mes_archivo]} {anio_archivo} de Supabase."
                        )
                        time.sleep(1)
                        st.rerun()
                    except Exception as e:
                        st.error(f"❌ No se pudo completar la limpieza mensual: {e}")
            else:
                st.success("✅ No existen registros para el mes seleccionado.")
        except Exception as e:
            st.error(f"❌ No se pudo consultar el mes seleccionado: {e}")

# =========================================================
# DESCARGA
# =========================================================
_, col_descarga_filtrados, _ = st.columns([1, 2, 1])
with col_descarga_filtrados:
    st.download_button(
        "📥 DESCARGAR FILTRADOS",
        data=excel_bytes(vista[COLS], None),
        file_name=f"SOLICITUDES_FILTRADAS_{date.today().isoformat()}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

st.markdown('<div class="footer"><strong>JC Control de Solicitudes — Almacén</strong><br>©JuanCarlosRamos - 2026 — Todos los derechos reservados</div>', unsafe_allow_html=True)
